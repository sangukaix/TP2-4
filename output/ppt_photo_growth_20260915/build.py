import copy
import json
import sys
from pathlib import Path
from unittest.mock import patch
from io import BytesIO
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from ai_server.app.proposal_presentation_v4 import create_strategy_proposal_presentation
from ai_server.app.proposal_case_outcomes import report_outcome

report=json.loads((ROOT/'output/final_validation_20260915/jeju_report.json').read_text(encoding='utf-8-sig'))
before=copy.deepcopy(report)
old=Presentation(ROOT/'output/final_validation_20260915/제주시_최종검증.pptx')
picture=next(s for s in old.slides[2].shapes if s.name=='visitor-photo')
source=next(s for s in report['evidence_sources'] if s.get('source_id')=='tour-api:1884191')
with patch('httpx.Client.get', side_effect=AssertionError('Offline export: network disabled')), patch(
    'ai_server.app.proposal_presentation_v4.download_images', return_value=[(source,BytesIO(picture.image.blob))]):
    result=create_strategy_proposal_presentation(report)
assert report==before, 'Saved report was mutated'
path=Path(__file__).parent/'제주시_기획서_디자인최종.pptx'
path.write_bytes(result.getvalue())
evidence=report_outcome(report)
(path.parent/'case_evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
print(path)

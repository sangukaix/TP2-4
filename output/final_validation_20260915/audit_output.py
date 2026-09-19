"""Read-only checks of this run's saved report and downloaded documents."""
import json
import math
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from docx import Document
from pptx import Presentation

folder = Path(__file__).resolve().parent
report = json.loads((folder / 'jeju_report.json').read_text(encoding='utf-8-sig'))
saved = json.loads((folder / 'jeju_saved_report.json').read_text(encoding='utf-8-sig'))
calculations = json.loads((folder / 'jeju_calculations.json').read_text(encoding='utf-8'))
scenario, demo = calculations['scenario'], calculations['demo']
basis = calculations['basis']
ppt = Presentation(folder / '제주시_최종검증.pptx')
doc = Document(folder / '제주시_최종검증.docx')
all_series = []
slides = []
for number, slide in enumerate(ppt.slides, 1):
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text)
        if shape.has_table:
            parts.extend(' | '.join(c.text for c in row.cells) for row in shape.table.rows)
        if shape.has_chart:
            for series in shape.chart.series:
                all_series.append({'slide': number, 'name': series.name, 'values': list(series.values)})
    slides.append({'slide': number, 'text': '\n'.join(parts)})
tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]
word_text = '\n'.join([p.text for p in doc.paragraphs] + [c for t in tables for r in t for c in r])
series_checks = {}
word_checks = {}
if scenario:
    for metric, scale in [('visitors', 10000), ('spending', 100000000)]:
        for kind in ['baseline', 'target']:
            key = kind + '_' + metric
            expected = [v / scale for v in scenario[key]]
            series_checks[key] = any(len(s['values']) == len(expected) and all(
                math.isclose(a, b, rel_tol=1e-9, abs_tol=1e-8)
                for a, b in zip(s['values'], expected)) for s in all_series)
            formatted = [f'{v:,.0f}' if metric == 'visitors' else f'{v/scale:,.2f}' for v in scenario[key]]
            word_checks[key] = {v: v in word_text for v in formatted}
estimate = report.get('reference_estimate') or {}
result = {
    'report_id': '14d7385fe7004f26898f89dbd9bc840c',
    'saved_report_matches': saved == report,
    'generation_mode': report.get('generation_mode'),
    'demo_forecast': demo,
    'title': (report.get('strategies') or [{}])[0].get('title'),
    'period': report.get('period'),
    'ppt_slides': len(ppt.slides),
    'word_tables': len(tables),
    'chart_checks': series_checks,
    'word_value_checks': word_checks,
    'estimate_sum_matches': sum(x['amount'] for x in estimate.get('items', [])) == estimate.get('total_krw'),
    'estimate_total': estimate.get('total_krw'),
    'totals': basis['totals'],
    'quality_review': report.get('quality_review'),
    'scenario': scenario,
    'every_page_visually_inspected': False,
}
(folder / 'jeju_document_audit.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
(folder / 'ppt_content.json').write_text(json.dumps(slides, ensure_ascii=False, indent=2), encoding='utf-8')
(folder / 'word_content.txt').write_text(word_text, encoding='utf-8')
print(json.dumps({k:v for k,v in result.items() if k not in ['quality_review','scenario']}, ensure_ascii=False, indent=2))

import copy
import json
import sys
from pathlib import Path
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT))
from ai_server.app.idea_proposal import prepare_idea_report
from ai_server.app.proposal_presentation_v3 import _scenario_rows

root=Path(__file__).parent
raw=json.loads((ROOT/'output/final_validation_20260915/jeju_report.json').read_text(encoding='utf-8-sig'))
original=copy.deepcopy(raw)
report=prepare_idea_report(raw)
assert raw==original
for key in ('ml_analysis','observed_findings','evidence_sources','quality_review','agent_trace'):
    assert report[key]==raw[key],key
assert prepare_idea_report(report)==report
plan=report['target_proposal_basis']['capacity_plan']
scenario=_scenario_rows(report)
assert abs(scenario['visitor_gap']-plan['central']['additional_visitors'])<.001
assert abs(scenario['spending_gap']-plan['central']['additional_spending_krw'])<.01
prs=Presentation(root/'deliverables/제주시_기획서_최종.pptx')
assert len(prs.slides)==17
assert len([s for p in prs.slides for s in p.shapes if s.has_chart])==6
assert len([s for p in prs.slides for s in p.shapes if s.has_table])==4
texts='\n'.join(s.text for p in prs.slides for s in p.shapes if s.has_text_frame)
assert '초기 목표 5%' not in texts
assert '2.3 운영 규모와 산출 근거' in texts
assert '2.4 기대 변화의 시나리오' in texts
assert '2,246' in texts and '1.57억 원' in texts
charts=[s.chart for s in prs.slides[5].shapes if s.has_chart]
assert abs(sum(charts[0].series[0].values)*10000-sum(scenario['baseline_visitors']))<1
assert abs(sum(charts[0].series[1].values)*10000-sum(scenario['target_visitors']))<1
assert report['reference_estimate']['total_krw']==168080000
result={'slides':17,'native_charts':6,'native_tables':4,'source_and_ml_unchanged':True,
        'central':plan['central'],'scenarios':plan['scenarios'],'estimate_krw':report['reference_estimate']['total_krw'],
        'population_adjusted':False,'learned_business_effect':False,
        'capacity_rate_sources':plan['rate_basis']}
(root/'audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
(root/'prepared_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print('PASS source preservation, shared targets, 17 slides, 6 native charts, 4 tables')

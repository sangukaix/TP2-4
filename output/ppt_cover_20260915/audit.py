from pathlib import Path
from hashlib import sha256
import json
from pptx import Presentation
r=Path(__file__).parent
oldr=r.parent/'ppt_spicus_20260915'
a=Presentation(oldr/'deliverables/제주시_기획서_최종.pptx');b=Presentation(r/'제주시_기획서_디자인최종.pptx')
changed=[p.name for p in (r/'render').glob('*.png') if sha256(p.read_bytes()).digest()!=sha256((oldr/'render'/p.name).read_bytes()).digest()]
assert changed==['slide-01.png','slide-03.png'],changed
photo=lambda sl,name:next(s.image.blob for s in sl.shapes if s.name==name)
assert sha256(photo(b.slides[0],'cover-regional-photo')).digest()!=sha256(photo(b.slides[2],'visitor-photo')).digest()
for x,y in zip(a.slides,b.slides):
    assert [[list(t.values) for t in s.chart.series] for s in x.shapes if s.has_chart]==[[list(t.values) for t in s.chart.series] for s in y.shapes if s.has_chart]
    assert [[c.text for rr in s.table.rows for c in rr.cells] for s in x.shapes if s.has_table]==[[c.text for rr in s.table.rows for c in rr.cells] for s in y.shapes if s.has_table]
for i,name in [(0,'cover-source-bg'),(2,'hero-source-bg')]:
    assert str(next(s.fill.fore_color.rgb for s in b.slides[i].shapes if s.name==name))=='0F315B'
from ai_server.app.proposal_cover_photo import candidates
report=json.loads((r.parent/'final_validation_20260915/jeju_report.json').read_text(encoding='utf-8-sig'))
s=candidates(report,'가마오름')
assert s and s[0]['title']=='곽지해수욕장'
assert all(x['address'].startswith(report['region_name']+' ') and x['title']!='가마오름' for x in s)
(r/'audit.json').write_text(json.dumps({'changed_slides':changed,'different_cover_photo':True,'blue_captions':True,'charts_tables_identical':True,'selected_photo':s[0]},ensure_ascii=False,indent=2),encoding='utf-8')
p=oldr/'finalize.mjs';(r/'finalize.mjs').write_text(p.read_text(encoding='utf-8').replace('ppt_spicus_20260915','ppt_cover_20260915'),encoding='utf-8')
print('Only slides 1/3 changed; distinct local photo, blue captions, numeric values preserved.')

import json
from pathlib import Path
from pptx import Presentation

root=Path(__file__).parent
previous=root.parent/'ppt_reference_cover_20260915'
old=Presentation(previous/'deliverables/제주시_기획서_최종.pptx')
new=Presentation(root/'deliverables/제주시_기획서_최종.pptx')
changed=[p.name for p in sorted((root/'render').glob('*.png'))
         if p.read_bytes()!=(previous/'render'/p.name).read_bytes()]
assert changed==[f'slide-{i:02}.png' for i in (2,7,8,9,10,12)],changed
def contents(prs):
    result=[]
    for i,slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if shape.has_text_frame:result.append((i,shape.name,shape.text))
            if shape.has_table:
                result.append((i,shape.name,[[c.text for c in row.cells] for row in shape.table.rows]))
            if shape.has_chart:
                result.append((i,shape.name,[(s.name,list(s.values)) for s in shape.chart.series]))
        result.append((i,'notes',slide.notes_slide.notes_text_frame.text))
    return result
assert contents(old)==contents(new),'Text, notes, charts or tables changed'
gradient=next(s for s in new.slides[1].shapes if s.name=='contents-rail')
assert gradient.fill.type==3
assert [str(s.color.rgb) for s in gradient.fill.gradient_stops]==['ECEEF1','BCC3CC']
pipeline=new.slides[11]
icon_count=sum(1 for s in pipeline.shapes if s.shape_type==13)
assert icon_count==8
result={'changed_slides':changed,'text_notes_charts_tables_unchanged':True,
        'gray_gradient_verified':True,'pipeline_icon_count':icon_count,
        'visual_review':'PowerPoint PNGs of all six changed slides reviewed'}
(root/'audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False))

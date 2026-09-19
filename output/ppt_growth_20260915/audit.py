import json,math,csv,sys
from pathlib import Path
sys.path.insert(0,str(Path.cwd()))
from ai_server.app.proposal_growth_comparison import comparison,rank
from ai_server.app.idea_proposal import prepare_idea_report
from pptx import Presentation
root=Path(__file__).resolve().parent
report=json.loads((root.parent/'final_validation_20260915/jeju_report.json').read_text(encoding='utf-8-sig'))
result=comparison(prepare_idea_report(report))
snap=json.loads((Path('artifacts/ml/comparison')/(result['signature']+'.json')).read_text(encoding='utf-8'))
assert rank([5,5,0,-3],5)==1 and rank([5,5,0,-3],0)==3
for key,short in [('visitors','visitors'),('spending_krw','spending')]:
 rates=[]
 for row in snap['rows']:
  m=row['metrics'][key]
  calculated=(sum(m['forecast'])-sum(m['previous']))/sum(m['previous'])*100
  assert math.isclose(calculated,m['growth_pct'],abs_tol=1e-10)
  rates.append(calculated)
 selected=result['metrics'][short]
 assert math.isclose(sum(rates)/len(rates),selected['national_mean_pct'],abs_tol=1e-10)
 assert 1+sum(v>selected['growth_pct']+1e-9 for v in rates)==selected['rank']
 prior=[v['metrics'][key]['growth_pct'] for v in snap['rows'] if v['region_code']!=result['region_code']]
 assert 1+sum(v>selected['target_growth_pct']+1e-9 for v in prior)==selected['target_rank']
old=Presentation(root.parent/'ppt_final_polish_20260915/deliverables/제주시_기획서_최종.pptx')
new=Presentation(root/'제주시_기획서_디자인최종.pptx')
def charts(p):return [[list(s.values) for s in sh.chart.series] for sl in p.slides for sh in sl.shapes if sh.has_chart]
def table(sl):return [[[c.text for c in r.cells] for r in sh.table.rows] for sh in sl.shapes if sh.has_table]
assert charts(old)==charts(new)
assert table(old.slides[9])==table(new.slides[9])
assert table(old.slides[10])==table(new.slides[10])
result['validation']={'regional_growth_rows_recalculated':len(snap['rows']),'chart_values_preserved':True,'kpi_budget_tables_preserved':True,'conditional_rank_and_ties_checked':True,'llm_calls':0,'model_training_calls':0}
(root/'audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
with (root/'regional_growth.csv').open('w',encoding='utf-8-sig',newline='') as f:
 writer=csv.DictWriter(f,fieldnames=list(result['cohort'][0]));writer.writeheader();writer.writerows(result['cohort'])
print(json.dumps(result['validation']))

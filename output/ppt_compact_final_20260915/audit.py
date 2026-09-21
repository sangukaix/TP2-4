import json
import sys
from pathlib import Path
from pptx import Presentation
from pptx.enum.text import PP_ALIGN

root=Path(__file__).parent
sys.path.insert(0,str(root.parents[1]))
from ai_server.app.proposal_theme_spicus import cover
from ai_server.app.proposal_layout_v7 import EMU
old=Presentation(root.parent/'ppt_balanced_colors_20260915/deliverables/제주시_기획서_최종.pptx')
new=Presentation(root/'deliverables/제주시_기획서_최종.pptx')
def numeric(prs):
    return [(i,s.name,[(x.name,list(x.values))for x in s.chart.series]) if s.has_chart
            else (i,s.name,[[c.text for c in row.cells]for row in s.table.rows])
            for i,slide in enumerate(prs.slides)for s in slide.shapes if s.has_chart or s.has_table]
assert numeric(old)==numeric(new)
def shape(slide,name):return next(s for s in slide.shapes if s.name==name)
titles=[shape(s,'title')for s in new.slides if any(x.name=='title'for x in s.shapes)]
assert all(p.alignment==PP_ALIGN.LEFT for s in titles for p in s.text_frame.paragraphs)
assert shape(new.slides[1],'toc-0-label').text=='개요'
assert shape(new.slides[2],'editorial-kicker').text=='0. 개요'
assert '출처' not in shape(new.slides[0],'cover-source-label').text
assert 'apis.data.go.kr/B551011/KorService2' in new.slides[0].notes_slide.notes_text_frame.text
assert all('지역 관광 전략기획안' not in s.text for s in new.slides[-1].shapes if s.has_text_frame)
sizes={}
for key,size in [('title',26),('case-title-0',14),('case-body-0',11),('case-image-caption-0',9)]:
    runs=[r for p in shape(new.slides[3],key).text_frame.paragraphs for r in p.runs]
    assert all(r.font.size.pt==size and r.font.name=='Century Gothic' for r in runs)
    assert all(r._r.xpath('./a:rPr/a:ea')[0].get('typeface')=='맑은 고딕'for r in runs)
    sizes[key]=size
# Check content stayed unchanged apart from explicitly requested labels and line wraps.
def normalize(value):return ' '.join(value.split())
allowed={(0,'cover-source-label'),(1,'contents-ko'),(1,'toc-0-label'),(2,'editorial-kicker')}
removed={(14,'cover-category'),(14,'cover-date')}
for i,(a,b) in enumerate(zip(old.slides,new.slides)):
    texts={s.name:s.text for s in b.shapes if s.has_text_frame and s.text}
    for s in a.shapes:
        if not s.has_text_frame or not s.text or (i,s.name) in allowed|removed:continue
        assert s.name in texts and normalize(s.text)==normalize(texts[s.name]),(i+1,s.name)
# Use the real caption renderer for short and long location names.
photo=shape(new.slides[0],'cover-regional-photo').image.blob
heights=[]
for location in ('제주시','제주특별자치도 제주시 애월읍 관광문화지구 해안 산책로와 지역 문화 공간'):
    p=Presentation();p.slide_width=1600*EMU;p.slide_height=900*EMU
    s=p.slides.add_slide(p.slide_layouts[6])
    cover(s,photo,region_name=location,photo_source={'title':'곽지해수욕장'})
    heights.append(shape(s,'cover-source-bg').height/EMU)
assert heights[0]<heights[1],heights
result={'native_charts_tables_unchanged':True,'numbered_headings_left_aligned':len(titles),
        'requested_text_changes_only':True,'reference_font_sizes_pt':sizes,
        'caption_short_long_height_px':heights,'cover_source_in_notes':True}
(root/'audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(result)

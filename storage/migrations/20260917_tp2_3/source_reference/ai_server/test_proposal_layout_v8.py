import unittest
from pptx import Presentation
from ai_server.test_proposal_presentation_v3_contract import _sample_report, _slide_text
from ai_server.app.proposal_presentation_v4 import create_strategy_proposal_presentation
from ai_server.app.proposal_layout_v8 import diverse_cases
from ai_server.app.agents.case_study_agent import _case_mechanism_family


class LayoutV8Test(unittest.TestCase):
    def test_diversity_uses_mechanisms_not_incidental_conditions(self):
        self.assertEqual(_case_mechanism_family({'title':'디지털 관광주민증','conditions':'숙박 가능'}),'return_visit')
        self.assertEqual(_case_mechanism_family({'title':'야간관광','conditions':'숙박시설 확보'}),'night_time_experience')
        report=_sample_report()
        report['evidence_sources']=[{'source_id':f'case:{i}','source_type':'benchmark_case','title':title}
            for i,title in enumerate(['야간관광','야간축제','야간 프로그램','반값 환급','관광주민증','숙박 할인'])]
        selected,_,_=diverse_cases(report)
        self.assertEqual(len({_case_mechanism_family(v) for v in selected}),4)

    def test_labels_totals_and_removed_sections(self):
        deck=Presentation(create_strategy_proposal_presentation(_sample_report()))
        slide=deck.slides[3]
        chart=next(s.chart for s in slide.shapes if s.has_chart)
        self.assertEqual(chart.series[0].name,'ML 기준 전망')
        self.assertGreater(chart.value_axis.minimum_scale,0)
        self.assertFalse(chart.has_legend)
        self.assertTrue(chart.plots[0].has_data_labels)
        self.assertTrue(all(series.data_labels.show_value for series in chart.series))
        legends=[s for s in slide.shapes if '-legend-label-' in s.name]
        self.assertEqual(len(legends),4)
        charts=[s for s in slide.shapes if s.has_chart]
        self.assertTrue(all(label.top>charts[0].top+charts[0].height for label in legends))
        bodies=[s for s in slide.shapes if s.name.startswith('block-body-')]
        self.assertEqual(len(bodies),3)
        self.assertEqual(len({r.font.size for s in bodies for p in s.text_frame.paragraphs for r in p.runs}),1)
        tables=[s.table for s in deck.slides[7].shapes if s.has_table]
        self.assertEqual(len(tables),2)
        self.assertTrue(all(t.cell(4,0).text=='3개월 월별 합계' for t in tables))
        self.assertIn('3개월 목표 추가 규모',_slide_text(deck.slides[7]))
        self.assertNotIn('시범사업 이후','\n'.join(_slide_text(s) for s in deck.slides))
        self.assertNotIn('FINAL',_slide_text(deck.slides[-1]))
        self.assertNotIn('TOUR INSIGHT',_slide_text(deck.slides[-1]))

if __name__=='__main__':unittest.main()

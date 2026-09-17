import unittest
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from ai_server.test_proposal_presentation_v3_contract import _sample_report, _slide_text
from ai_server.app.proposal_presentation_v4 import create_strategy_proposal_presentation
from ai_server.app.proposal_layout_v9 import execution_copy, case_selection_basis, execution_groups, period

class ExecutionDetailTest(unittest.TestCase):
    def test_five_steps_preserve_operation_and_evaluation_dates_in_both_outputs(self):
        from copy import deepcopy
        from docx import Document
        from ai_server.app.proposal_document import create_strategy_proposal_document
        r = _sample_report()
        r['strategies'][0]['implementation_steps'] = [
            {'step':1, 'schedule':'2026-10', 'task':'참여 시설 조사', 'deliverable':'시설 목록'},
            {'step':2, 'schedule':'2026-10', 'task':'콘텐츠와 안전 계획 수립', 'deliverable':'운영 계획'},
            {'step':3, 'schedule':'2026-10', 'task':'예약 시스템 점검', 'deliverable':'점검 결과'},
            {'step':4, 'schedule':'2026-11 ~ 2026-12', 'task':'프로그램 시범 운영', 'deliverable':'운영 원장'},
            {'step':5, 'schedule':'2026-12', 'task':'성과 분석', 'deliverable':'결과 보고서'},
        ]
        before = deepcopy(r)
        groups = execution_groups(r)
        self.assertEqual([v for g in groups for v in g], r['strategies'][0]['implementation_steps'])
        self.assertEqual(period(groups[2]), '2026-11 ~ 2026-12')
        self.assertEqual(period(groups[3]), '2026-12')
        deck = Presentation(create_strategy_proposal_presentation(r))
        periods = {s.name:s.text for s in deck.slides[6].shapes if s.name.startswith('execution-period-')}
        self.assertIn('2026-11 ~ 2026-12', periods['execution-period-2'])
        self.assertTrue(periods['execution-period-3'].startswith('2026-12'))
        doc = Document(create_strategy_proposal_document(r))
        paragraphs = [p.text for p in doc.paragraphs]
        pilot = paragraphs.index('03  시범 운영')
        self.assertIn('2026-11 ~ 2026-12', paragraphs[pilot+1])
        self.assertEqual(r, before)

    def test_selection_uses_linked_case_without_environment_claim(self):
        r=_sample_report()
        r['strategies'][0]['title']='여행비 환급'
        r['strategies'][0]['solution']='지역상품권 환급'
        r['planning_decision']={'selected_candidate_id':'chosen','design_candidates':[
            {'candidate_id':'chosen','title':'관광 프로그램','case_source_ids':['case:chosen'],'mechanism':'지역 상점 재이용'}]}
        r['evidence_sources']=[
            {'source_id':'case:other','source_type':'benchmark_case','title':'다른 지역'},
            {'source_id':'case:chosen','source_type':'benchmark_case','title':'선정 사례','operating_model':'지역상품권으로 돌려받아 다시 이용',
             'source_url':'https://example.go.kr/case'}]
        basis=case_selection_basis(r)
        self.assertEqual(basis['source']['source_id'],'case:chosen')
        deck=Presentation(create_strategy_proposal_presentation(r))
        slide=deck.slides[5]
        matrix=next(s for s in slide.shapes if s.name=='selection-matrix')
        self.assertEqual(len(matrix.table.rows),4)
        self.assertNotIn('지리·교통·관광자원', ' '.join(c.text for row in matrix.table.rows for c in row.cells))
        self.assertIn('https://example.go.kr/case',slide.notes_slide.notes_text_frame.text)
        self.assertNotIn('다른 지역',_slide_text(slide))
    def test_page_order_and_overview_headers(self):
        p=Presentation(create_strategy_proposal_presentation(_sample_report()))
        self.assertIn('지역 적용 방법',_slide_text(p.slides[5]))
        self.assertIn('4단계 실행 가이드 예시안',_slide_text(p.slides[6]))
        self.assertIn('머신러닝 예측값',_slide_text(p.slides[7]))
        self.assertEqual(len([s for s in p.slides[6].shapes if s.name.startswith('flow-node-')]),4)
    def test_refund_instructions_not_for_unrelated_program(self):
        r=_sample_report();r['strategies'][0]['title']='반값 여행 환급'
        self.assertIn('영수증',str(execution_copy(r)))
        r['strategies'][0]['title']='지역 문화 체험';r['strategies'][0]['solution']='문화시설 예약 운영';r['planning_decision']={}
        self.assertNotIn('환급',str(execution_copy(r)))
if __name__=='__main__':unittest.main()

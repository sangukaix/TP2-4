from copy import deepcopy
import unittest
from ai_server.app.proposal_evidence import saved_budget_view


class SavedBudgetViewTest(unittest.TestCase):
    def test_native_ppt_table_and_notes_use_saved_budget(self):
        from pptx import Presentation
        from ai_server.test_proposal_presentation_v3_contract import _sample_report
        from ai_server.app.proposal_presentation import create_strategy_proposal_presentation
        report = _sample_report()
        report['generation_mode'] = 'openai'
        report['strategies'][0]['budget'] = '총 1,000,000,000원 (기획 가정: 항목별 단가 미확정)'
        deck = Presentation(create_strategy_proposal_presentation(report))
        text = '\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
        cells = '\n'.join(cell.text for slide in deck.slides for shape in slide.shapes if shape.has_table
                          for row in shape.table.rows for cell in row.cells)
        self.assertIn('참여 목표에 따른 예상 사업비', text)
        self.assertIn('예상 견적이며 실제 액수와 다를 수 있습니다', text)
        self.assertEqual(report['strategies'][0]['budget'], '총 1,000,000,000원 (기획 가정: 항목별 단가 미확정)')
        self.assertIn('기존 수단 설정 1식 × 가정 단가 8,000,000원', cells)
        notes = '\n'.join(slide.notes_slide.notes_text_frame.text for slide in deck.slides)
        self.assertNotIn('reference-estimate-v1', notes)

    def test_report_total_is_not_replaced_by_generic_pilot(self):
        report = {'strategies': [{'budget': '총 1,000,000,000원 (기획 가정: 세부 항목 합계)'}]}
        before = deepcopy(report)
        view = saved_budget_view(report)
        self.assertEqual(view['total_label'], '총 1,000,000,000원')
        self.assertEqual(view['rows'][0][1], report['strategies'][0]['budget'])
        self.assertEqual(report, before)
        self.assertEqual(view['sources'], [])

    def test_formula_without_total_does_not_create_money(self):
        view = saved_budget_view({'strategies': [{'budget': '운영일수 × 일 단가 + 인원 × 수당'}]})
        self.assertEqual(view['total_label'], '총액 미확정')
        self.assertIn('운영일수 × 일 단가', view['original_text'])

    def test_multiple_totals_are_not_silently_selected(self):
        view = saved_budget_view({'strategies': [{'budget': '안1 총 10억원; 안2 총 5억원'}]})
        self.assertEqual(view['total_label'], '총액 미확정')
        self.assertEqual(len(view['rows']), 2)

    def test_empty_budget_is_explicit_not_a_default_quote(self):
        self.assertEqual(saved_budget_view({})['rows'][0][2], '미정')

    def test_all_lines_preserved_when_row_count_is_bounded(self):
        text = '\n'.join(f'항목 {i} × 단가 {i}' for i in range(12))
        view = saved_budget_view({'strategies': [{'budget': text}]})
        self.assertEqual(len(view['rows']), 6)
        joined = ' '.join(row[1] for row in view['rows'])
        for i in range(12):
            self.assertIn(f'항목 {i} × 단가 {i}', joined)

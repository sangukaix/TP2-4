"""New layout data and text invariants, without network or paid LLM calls."""
import copy
import unittest
from pptx import Presentation
from ai_server.test_proposal_presentation_v3_contract import _sample_report
from ai_server.app.proposal_presentation_v4 import create_strategy_proposal_presentation
from ai_server.app.proposal_layout_v7 import wrap_words
from ai_server.presentation_test_support import slide_with_title

class LayoutV7Test(unittest.TestCase):
    def test_native_chart_values_and_input_immutability(self):
        report=_sample_report();before=copy.deepcopy(report)
        deck=Presentation(create_strategy_proposal_presentation(report))
        charts=[s.chart for s in slide_with_title(deck,'2.1 사업 목표').shapes if s.has_chart]
        self.assertEqual(list(charts[0].series[0].values),[1800,1810,1820])
        self.assertEqual(list(charts[1].series[0].values),[4700,4740,4780])
        self.assertEqual(len(charts[0].series),2)
        self.assertAlmostEqual(charts[0].series[1].values[-1],1820*1.01)
        self.assertEqual(report,before)

    def test_automatic_targets_follow_capacity_not_fixed_five_percent(self):
        from ai_server.app.idea_proposal import prepare_idea_report
        report=_sample_report();report['execution_scenario']=None
        deck=Presentation(create_strategy_proposal_presentation(report))
        charts=[s.chart for s in slide_with_title(deck,'2.1 사업 목표').shapes if s.has_chart]
        self.assertEqual(len(charts),2)
        self.assertTrue(all(len(chart.series)==2 for chart in charts))
        table=next(s.table for s in slide_with_title(deck,'3.3 머신러닝 예측값과 목표 KPI').shapes if s.has_table)
        self.assertNotEqual(table.cell(1,2).text,'목표율 미입력')
        prepared=prepare_idea_report(report)
        self.assertEqual(prepared['execution_scenario']['target_origin'],'operating_capacity')
        pct=prepared['execution_scenario']['visitor_target_pct']
        self.assertNotEqual(pct,5)
        self.assertAlmostEqual(charts[0].series[1].values[-1],charts[0].series[0].values[-1]*(1+pct/100))
        self.assertIsNone(report['execution_scenario'])

    def test_five_steps_are_kept_in_four_rows(self):
        deck=Presentation(create_strategy_proposal_presentation(_sample_report()))
        slide=slide_with_title(deck,'3.2 4단계 실행 가이드 예시안')
        self.assertEqual(len([s for s in slide.shapes if s.name.startswith('flow-node-')]),4)
        self.assertIn('실행 작업 5',slide.notes_slide.notes_text_frame.text)
        self.assertIn('확인 산출물 5',slide.notes_slide.notes_text_frame.text)

    def test_wrapping_preserves_words_without_last_word_orphan(self):
        sentence='방문객이 지역 내에서 지출한 금액의 일부를 지역상품권으로 돌려받는 방식으로 혜택이 다시 지역 내 상점에서 사용되도록 설계'
        lines=wrap_words(sentence,883,34)
        self.assertEqual(' '.join(lines),sentence)
        self.assertGreater(len(lines[-1].split()),1)
        self.assertTrue(all(word in sentence.split() for line in lines for word in line.split()))

if __name__=='__main__':unittest.main()

import unittest
from pptx import Presentation
from ai_server.app.case_images import case_image, IMAGES
from ai_server.app.proposal_layout_v10 import case_cards


class CaseImageTest(unittest.TestCase):
    def test_unknown_case_labels_reference_and_preserves_real_location(self):
        result=case_image({'source_id':'case:unknown','case_region':'전주시','intervention':'야간 관광'})
        self.assertEqual(result['match_kind'],'similar_operation')
        self.assertIn('다른 지역 참고 사진',result['caption'])
        self.assertIn('여수',result['caption'])
        self.assertNotIn('전주',result['caption'])
        self.assertTrue(result['path'].is_file())
        fallback=case_image({'source_id':'case:unknown'})
        self.assertEqual(fallback['match_kind'],'generated_operating_example')
        self.assertIn('AI 생성 이미지',fallback['credit'])
        self.assertTrue(fallback['path'].is_file())
        self.assertTrue(all(case_image({'source_id':key}) for key in IMAGES))

    def test_card_uses_official_asset_and_credit_without_generated_fallback(self):
        p=Presentation();s=p.slides.add_slide(p.slide_layouts[6])
        case_cards(s,{'region_name':'서울특별시','strategies':[{'solution':'여행비 환급'}],
            'evidence_sources':[{'source_id':'case:gangjin_half_price_2024','source_type':'benchmark_case',
                'source_url':'https://www.gangjintour.com/','operating_model':'상품권 환급','intervention':'반값 환급'}]})
        pictures=[sh for sh in s.shapes if sh.name=='case-photo-0']
        self.assertEqual(len(pictures),1)
        self.assertEqual(pictures[0].crop_left,0)
        text=' '.join(sh.text for sh in s.shapes if sh.has_text_frame)
        self.assertIn('공식 홈페이지',text)
        self.assertNotIn('OpenAI',text)
        self.assertNotIn('tourism_operations',s.notes_slide.notes_text_frame.text)

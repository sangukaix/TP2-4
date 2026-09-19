import unittest
from pptx import Presentation
from ai_server.app.proposal_business_images import select_business_image

class BusinessImageTest(unittest.TestCase):
    def test_themes_and_selected_mechanism(self):
        for words, filename in [('쿠폰 환급','refund.png'),('야간 공연','night.png'),('숙박 상품','stay.png'),('관광 교통','mobility.png'),('공예 체험','experience.png'),('예약 입장','reservation.png')]:
            r={'strategies':[{'solution':words}]}
            image=select_business_image(r,Presentation())
            self.assertEqual(image['path'].name,filename)
            self.assertIn('AI 생성',image['caption'])
        r={'strategies':[{'solution':'쿠폰 환급'}],'planning_decision':{'selected_candidate_id':'a','design_candidates':[{'candidate_id':'a','mechanism':'야간 공연'}]}}
        self.assertEqual(select_business_image(r,Presentation())['family'],'night_time_experience')
    def test_duplicate_content_is_rejected(self):
        p=Presentation();r={'strategies':[{'solution':'쿠폰 환급'}]}
        info=select_business_image(r,p)
        p.slides.add_slide(p.slide_layouts[6]).shapes.add_picture(str(info['path']),0,0)
        with self.assertRaises(ValueError):select_business_image(r,p)

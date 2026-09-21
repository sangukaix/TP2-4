"""Exact ML chart bindings, no invented uplift, and complete source pagination."""
from copy import deepcopy
import unittest
from pptx import Presentation
from ai_server.test_proposal_presentation_v3_contract import _sample_report, _slide_text
from ai_server.app.proposal_presentation import create_strategy_proposal_presentation
from ai_server.app.proposal_evidence import build_reference_estimate, complete_source_records, source_display_name
from ai_server.app.report_projection import select_report_forecast
from ai_server.presentation_test_support import slide_with_title


def table_text(slide):
    return '\n'.join(cell.text for s in slide.shapes if getattr(s,'has_table',False)
                     for row in s.table.rows for cell in row.cells)


class ProposalEvidenceTest(unittest.TestCase):
    def test_exact_months_and_values_are_used_without_observation_splice(self):
        report=_sample_report()
        original=deepcopy(report)
        deck=Presentation(create_strategy_proposal_presentation(report))
        rows=select_report_forecast(report)['rows']
        charts=[s.chart for s in slide_with_title(deck,'3.1 사업 목표').shapes if getattr(s,'has_chart',False)]
        self.assertEqual(len(charts),2)
        for chart,key,unit in zip(charts,['visitors','spending_krw'],[1e4,1e8]):
            self.assertEqual(len(chart.plots[0].categories),len(rows))
            for actual,row in zip(chart.series[0].values,rows):self.assertAlmostEqual(actual*unit,row[key],delta=.1)
            self.assertLessEqual(chart.value_axis.minimum_scale,min(chart.series[0].values))
            self.assertGreaterEqual(chart.value_axis.maximum_scale,max(chart.series[0].values))
        self.assertEqual(report,original)

    def test_target_uses_final_month_not_first_month(self):
        report=_sample_report();deck=Presentation(create_strategy_proposal_presentation(report))
        rows=select_report_forecast(report)['rows']; last=rows[-1]
        data=table_text(slide_with_title(deck,'4.2 머신러닝 예측값과 목표 KPI'))
        self.assertIn(f"{last['visitors']*1.01:,.0f}명",data)
        self.assertIn(f"{last['spending_krw']*1.02/1e8:,.2f}억 원",data)

    def test_missing_target_does_not_invent_regional_uplift(self):
        report=_sample_report();report['execution_scenario']=None
        report['strategies'][0]['title']='지역 반값 여행 환급'
        deck=Presentation(create_strategy_proposal_presentation(report))
        slide=slide_with_title(deck,'4.2 머신러닝 예측값과 목표 KPI');text=_slide_text(slide)+table_text(slide)
        self.assertIn('목표 KPI',text);self.assertNotIn('목표율 미입력',text)
        self.assertNotIn('산출 보류',text)
        self.assertNotIn('운영 후 실제 확인',text)
        self.assertIsNone(report['execution_scenario'])

    def test_reference_budget_sums_and_hard_limit(self):
        report=_sample_report();report['strategies'][0]['title']='반값 환급 여행'
        estimate=build_reference_estimate(report)
        self.assertGreater(estimate['quantity'],1000)
        larger=__import__('copy').deepcopy(report)
        larger['execution_scenario']={'visitor_target_pct':2,'spending_target_pct':2}
        # Explicit user targets do not invent a larger operating capacity/cost.
        self.assertEqual(build_reference_estimate(larger)['total_krw'],estimate['total_krw'])
        self.assertEqual(sum(i['amount'] for i in estimate['items']),estimate['total_krw'])
        report['planning_brief'].update(budget_hard_limit=True,budget_max_krw=40_000_000)
        estimate=build_reference_estimate(report)
        self.assertLessEqual(estimate['total_krw'],40_000_000)
        report['planning_brief']['budget_max_krw']=1
        self.assertFalse(build_reference_estimate(report)['within_hard_budget'])

    def test_all_sources_survive_pagination_and_notes(self):
        report=_sample_report()
        for n in range(45):
            report['evidence_sources'].append({'source_id':f'data:{n}','source_type':'dataset',
              'title':f'지역별 자료 {n:02d} 이동통신 방문 특성과 소비 현황 월별 원자료',
              'source_url':f'https://example.go.kr/source/{n}'})
        deck=Presentation(create_strategy_proposal_presentation(report))
        text='\n'.join(_slide_text(s) for s in list(deck.slides)[9:])
        notes='\n'.join(slide.notes_slide.notes_text_frame.text for slide in deck.slides)
        self.assertGreaterEqual(len(deck.slides),12)
        for n in range(45):
            self.assertIn(f'지역별 자료 {n:02d}',text)
            self.assertIn(f'https://example.go.kr/source/{n}',notes)
        self.assertNotIn('…',text)
        self.assertIn('감사합니다',_slide_text(deck.slides[-1]))
        self.assertNotIn('AI 기획·검수',text)

    def test_filename_keeps_metric_folder_not_just_zip_date(self):
        name=source_display_name({'source_file_name':'원주시/숙박_체류시간/2026_01_06.zip'})
        self.assertEqual(name,'숙박_체류시간/2026_01_06.zip')

    def test_pdf_source_listing_keeps_all_selected_page_references(self):
        records=complete_source_records({'evidence_sources':[
            {'source_id':'pdf:1','chunk_id':'pdf:1:p34','page_references':'PDF 34쪽','title':'공식 PDF'},
            {'source_id':'pdf:1','chunk_id':'pdf:1:p35','page_references':'PDF 35쪽','title':'공식 PDF'},
        ]})
        self.assertEqual(len(records),1)
        self.assertEqual(records[0]['page_references'],'PDF 34쪽, PDF 35쪽')
        self.assertEqual(records[0]['source_chunk_count'],2)

    def test_observed_metric_units_and_missing_summary_fallback(self):
        from ai_server.app.proposal_presentation_v4 import _metric_finding, _short_metric, _card_value
        report = {'observed_findings': [{'metric': '외지인 방문자 증감률', 'value': '-2.7%'}],
                  'monthly_trend': [{'month': '2026.06', 'visitors': 2285734, 'spending_krw': 60562672000},
                                    {'month': '2026.07', 'visitors': 9999999, 'is_forecast': True}]}
        self.assertEqual(_metric_finding(report)['value'], '2,285,734명')
        self.assertEqual(_metric_finding(report, spending=True)['value'], '60,562,672,000원')
        self.assertEqual(_short_metric('외지인 방문자 증감률'), '방문자 증감률')
        self.assertEqual(_card_value('관광소비액', '4,819억 원'), '4,819억 원')
        self.assertEqual(_card_value('소비 비중', '84.2%'), '84%')


if __name__=='__main__':unittest.main()

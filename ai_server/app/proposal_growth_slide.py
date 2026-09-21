"""Opening purpose page: comparable growth, not population-adjusted impact."""
import json
import logging
from .proposal_layout_v7 import header, text, rect, BLUE, CYAN, ORANGE, SLATE, WHITE, PALE, INK
from .proposal_layout_v7 import clear, EMU


def growth_bars(slide, name, region, row, y):
    """Editable signed bars share a real zero; values are percentage points."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_TICK_MARK, XL_TICK_LABEL_POSITION
    from pptx.dml.color import RGBColor
    from pptx.util import Pt
    from pptx.oxml.xmlchemy import OxmlElement
    # Use one six-decimal display value for both OOXML cache and Excel cells.
    # The full calculation precision remains in the source/slide notes.
    values = [round(row['growth_pct'], 6), round(row['national_mean_pct'], 6)]
    data = CategoryChartData()
    data.categories = [region, '전국 비교 평균']
    data.add_series('전년 동기 대비 증가율', values)
    shape = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, 561*EMU, y*EMU, 910*EMU, 150*EMU, data)
    shape.name = name
    chart = shape.chart
    chart.has_legend = False
    chart.has_title = False
    chart.chart_style = 2
    plot = chart.plots[0]
    plot.series[0].invert_if_negative = False
    plot.gap_width = 75
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
    labels.number_format = '+0.00"%";-0.00"%";0.00"%"'
    labels.font.name = 'Noto Sans KR'
    labels.font.size = Pt(16)
    labels.font.bold = True
    labels.font.color.rgb = RGBColor.from_string(BLUE)
    for point, color in zip(plot.series[0].points, (BLUE, CYAN)):
        invert = OxmlElement('c:invertIfNegative')
        invert.set('val', '0')
        point._ser.get_or_add_dPt_for_point(point._idx).insert(1, invert)
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = RGBColor.from_string(color)
        point.format.line.fill.background()
    axis = chart.value_axis
    span = max(max(values)-min(values), max(abs(v) for v in values), 1)
    axis.minimum_scale = min(0, min(values))-span*.28
    axis.maximum_scale = max(0, max(values))+span*.28
    axis.has_major_gridlines = False
    axis.tick_label_position = XL_TICK_LABEL_POSITION.NONE
    axis.major_tick_mark = axis.minor_tick_mark = XL_TICK_MARK.NONE
    axis.format.line.fill.background()
    cat = chart.category_axis
    cat.reverse_order = True
    cat.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    cat.major_tick_mark = cat.minor_tick_mark = XL_TICK_MARK.NONE
    cat.tick_labels.font.name = 'Noto Sans KR'
    cat.tick_labels.font.size = Pt(15)
    cat.tick_labels.font.color.rgb = RGBColor.from_string(SLATE)
    cat.format.line.color.rgb = RGBColor.from_string('C8D7E3')
    # pptx's bar defaults use negative axis IDs; OOXML axis IDs are unsigned.
    for element in chart._chartSpace.xpath('.//c:axId | .//c:crossAx'):
        element.set('val', str(abs(int(element.get('val')))))
    return shape


def growth_page(slide, report):
    from .proposal_growth_comparison import comparison
    from .proposal_presentation_v4 import _scenario_for_display
    region = str(report.get('region_name','')).split()[-1]
    # Keep the original full-height regional photo and its linked caption.
    clear(slide, keep=('visitor-photo', 'hero-source-bg', 'hero-source-label'))
    text(slide,'title',f'{region} 관광 발전 기획안',561,86,950,77,42,BLUE,True)
    rect(slide,'title-rule',561,181,933,1,'D4DFE8')
    scenario, _ = _scenario_for_display(report)
    try:
        result = comparison(report)
    except (OSError, ValueError, KeyError) as exc:
        logging.getLogger(__name__).warning('Growth comparison unavailable: %s', type(exc).__name__)
        result = None
    if not result:
        text(slide,'purpose','월별 관광 흐름을 확인하고, 공식 사례에서 지역에 맞는 실행 아이디어를 찾습니다.',561,213,933,100,29,SLATE)
        if scenario:
            from .proposal_layout_v7 import chart
            chart(slide,'purpose-visitors',561,345,440,345,scenario,'visitors',10000)
            chart(slide,'purpose-spending',1030,345,440,345,scenario,'spending',1e8)
            text(slide,'purpose-key','방문·소비 전망 → 공식 사례 비교 → 지역에 맞춘 실행 제안',561,770,933,80,26,BLUE,True)
        else:
            text(slide,'purpose-key','지역 관측 지표와 공식 사례의 운영 방식을 연결하여 실행 방향을 제안합니다.',561,365,933,140,32,BLUE,True)
        slide.notes_slide.notes_text_frame.text='동일 기간·모델 지문이 맞는 전국 비교가 없을 때의 지역 전망 구성. 인구 보정·전국 순위를 임의로 만들지 않음.'
        return
    start,end=result['months'][0],result['months'][-1]
    text(slide,'purpose-period',f"향후 3개월 전망 · {start[:4]}년 {start[4:]}–{end[4:]}월",561,211,933,42,26,SLATE)
    evaluation = ((report.get('ml_analysis') or {}).get('evaluation') or {}).get('metrics') or {}
    seasonal_baseline = any(
        abs(result['metrics'][key]['growth_pct']) < 1e-9
        and (evaluation.get(metric) or {}).get('selected_model') == 'seasonal_naive_previous_year_same_month'
        for key, metric in (('visitors', 'visitors'), ('spending', 'spending_krw'))
    )
    if seasonal_baseline:
        comparison_copy = ('전년 같은 기간 대비 증가율입니다.\n'
                           '0%는 데이터 없음이 아니라, 검증 결과 전년 동월 계절 기준선이 선택된 전망입니다.')
        text(slide,'purpose-comparison',comparison_copy,561,246,933,60,18,SLATE)
    else:
        text(slide,'purpose-comparison','전년 같은 기간 대비 증가율을 전국 비교 평균과 비교합니다.',561,253,933,40,22,SLATE)
    for i,(key,label) in enumerate([('visitors','관광객 방문자 수'),('spending','관광소비액')]):
        row=result['metrics'][key]; y=315+i*245
        text(slide,f'growth-label-{i}',label,561,y,400,40,27,BLUE,True)
        total=sum(row['forecast'])
        amount=f'{total/10000:,.1f}만 명' if key=='visitors' else f'{total/1e8:,.1f}억 원'
        text(slide,f'growth-total-{i}',f'3개월 예측 합계  {amount}',1000,y+4,494,37,21,SLATE)
        growth_bars(slide,f'growth-chart-{i}',region,row,y+41)
        relation='높습니다' if row['gap_pp']>=0 else '낮습니다'
        text(slide,f'growth-gap-{i}',f"{region}의 예상 증가율은 전국 비교 평균보다 {abs(row['gap_pp']):.2f}%p {relation}.",561,y+189,933,44,23,BLUE,True)
    text(slide,'purpose-note',f"비교 대상 {result['count']}개 지역 · 한국관광 데이터랩 {result['origin'][:4]}.{result['origin'][4:]}까지의 관측값과 저장 ML 모델\n비교 평균 = 지역별 전년 동기 대비 증가율의 단순평균. 인구 보정 전 탐색 전망입니다.",561,833,933,62,16,SLATE)
    slide.notes_slide.notes_text_frame.text=json.dumps({
        'comparison':result,'method':'Saved region-specific models; same cutoff, same months; no retraining. Equal-weight mean of regional quarterly YoY growth. Competition ranks.',
        'source':'한국관광 데이터랩 공식 다운로드; data/processed/ml/<region_code>/monthly_demand.csv; artifacts/ml/<region_code>/demand_model.joblib',
        'limits':'No resident population or density adjustment. Forecast distance from last observed month may exceed 3 months. Rank changes are conditional scenarios, not causal estimates.'},ensure_ascii=False,indent=2)


SECTIONS = {
    '지역별 참고 사례':('1. 지역별 사례','1.1 지역별 참고 사례'),
    '사례 실적':('1. 지역별 사례','1.2 적용 사례 운영 방식'),
    '참고 사례 운영 방식':('1. 지역별 사례','1.2 적용 사례 운영 방식'),
    '적용 사례 운영 방식':('1. 지역별 사례','1.2 적용 사례 운영 방식'),
    '사업 목표':('2. 사업 목표','2.1 사업 목표'),
    '목표 KPI 산출근거':('2. 사업 목표','2.2 목표 KPI 산출근거'),
    '이 사업을 참고한 이유와 지역 적용 방법':('3. 지역 적용과 실행','3.1 사례 선정과 지역 적용'),
    '4단계 실행 가이드 예시안':('3. 지역 적용과 실행','3.2 4단계 실행 가이드 예시안'),
    '머신러닝 예측값과 목표 KPI':('3. 지역 적용과 실행','3.3 머신러닝 예측값과 목표 KPI'),
    '견적 예시안':('4. 예상 견적','4.1 견적 예시안'),
    '기획서 생성 파이프라인':('5. 생성 과정','5.1 기획서 생성 파이프라인'),
    '근거·데이터·출처':('6. 근거와 출처','6.1 근거·데이터·출처'),
}


def section_labels(prs):
    from .proposal_visual_finish import restyle
    from .proposal_layout_v7 import EMU
    for i,slide in enumerate(prs.slides):
        title=next((s for s in slide.shapes if s.name=='title' and s.has_text_frame),None)
        if not title: continue
        raw=title.text
        main,sub=SECTIONS.get(raw,('0. 개요',raw) if i==2 else ('5. 생성 과정',raw))
        if raw=='근거·데이터·출처' and i>0 and any(s.name=='title' and s.text.startswith('6.1') for s in prs.slides[i-1].shapes if s.has_text_frame):
            sub='6.2 근거·데이터·출처 계속'
        title.text=sub; restyle(title,42)
        kicker=next((s for s in slide.shapes if s.name=='editorial-kicker'),None)
        if kicker:
            kicker.text=main;kicker.height=35*EMU;restyle(kicker,23,'008795',True)
        else:
            text(slide,'editorial-kicker',main,106,38,1150,37,23,'008795',True)

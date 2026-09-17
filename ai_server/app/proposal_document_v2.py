"""A4 proposal, sharing the approved PPT's data and narrative projections."""
from io import BytesIO
import re

from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from .idea_proposal import prepare_idea_report
from .proposal_calculation_basis import calculation_basis
from .proposal_presentation_v4 import _scenario_for_display, _public_copy
from .proposal_evidence import complete_source_records, source_display_name
from .case_recommendation import report_cases
from .case_images import case_image
from .proposal_case_outcomes import report_outcome
from .proposal_layout_v9 import execution_groups, execution_copy, period, LABELS, ROLES

BLUE = '0054A6'
GREEN = '16845B'
GRAY = '536477'


def prose(value):
    value = re.sub(r'\s*\(?\[[^\]]+\]\(https?://[^\s]+\)\)?', '', _public_copy(value))
    return value.strip()


def paragraph(parent, value='', *, size=10, bold=False, color='263445', style=None):
    p = parent.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.18
    p.paragraph_format.widow_control = True
    snap = OxmlElement('w:snapToGrid'); snap.set(qn('w:val'), '0')
    p._p.get_or_add_pPr().append(snap)
    r = p.add_run(str(value))
    r.font.name = '맑은 고딕'
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    return p


def heading(doc, value, new_page=False):
    p = paragraph(doc, value, size=18, bold=True, color=BLUE, style='Heading 1')
    p.paragraph_format.page_break_before = new_page
    p.paragraph_format.space_after = Pt(15)
    return p


def subheading(doc, value):
    p = paragraph(doc, value, size=12, bold=True, color=BLUE, style='Heading 2')
    p.paragraph_format.space_before = Pt(13)
    return p


def source_link(doc, url, label='공식 원문 보기', size=8):
    """Retain the complete URL in a clickable relation without printing query noise."""
    if not url:
        return
    p = paragraph(doc, '', size=size)
    p.paragraph_format.space_after = Pt(6)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), doc.part.relate_to(str(url), RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), BLUE); props.append(color)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size*2))); props.append(sz)
    fonts = OxmlElement('w:rFonts')
    for attr in ('ascii', 'hAnsi', 'eastAsia'): fonts.set(qn('w:'+attr), '맑은 고딕')
    props.append(fonts); run.append(props)
    text = OxmlElement('w:t'); text.text = label; run.append(text)
    link.append(run); p._p.append(link)
    return p


def table(doc, rows, widths, green_last=False):
    t = doc.add_table(rows=0, cols=len(widths))
    t.autofit = False
    for column, width in zip(t.columns, widths):
        column.width = Mm(width)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        element = OxmlElement('w:' + edge)
        for key, val in [('val', 'single'), ('sz', '4'), ('color', 'FFFFFF')]:
            element.set(qn('w:' + key), val)
        borders.append(element)
    t._tbl.tblPr.append(borders)
    for i, values in enumerate(rows):
        row = t.add_row()
        if i == 0:
            row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for j, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Mm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), (GREEN if green_last and j == len(widths)-1 else BLUE) if i == 0 else ('EFF5F8' if i % 2 == 0 else 'FFFFFF'))
            props.append(shade)
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'bottom', 'left', 'right'):
                el = OxmlElement('w:' + side); el.set(qn('w:w'), '100'); el.set(qn('w:type'), 'dxa'); margins.append(el)
            props.append(margins)
            p = paragraph(cell, value, size=9, bold=i == 0, color='FFFFFF' if i == 0 else '263445')
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            if i > 0 and j > 0 and re.fullmatch(r'[\d,.]+(?:원|명|억 원|건|%|일|분)', str(value)):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if i > 0 and str(value).startswith('↑'):
                p.runs[0].font.color.rgb = RGBColor.from_string('E76700')
                p.runs[0].font.bold = True
            if i > 0 and str(values[0]) in ('합계', '월별 합계'):
                p.runs[0].font.bold = True
            cell._tc.remove(cell.paragraphs[0]._p)
    paragraph(doc, '', size=2).paragraph_format.space_after = Pt(2)
    return t


def forecast_chart(scenario, metric=None):
    # The Word plot consumes exactly the arrays that drive the native PPT charts.
    from .proposal_document import plt
    specs = [('visitors', '방문자 수 · 만 명', 10000), ('spending', '관광소비액 · 억 원', 100000000)]
    if metric: specs = [v for v in specs if v[0] == metric]
    fig, axes = plt.subplots(1, len(specs), figsize=(8.1, 3.3), dpi=190, squeeze=False)
    for ax, (key, label, unit) in zip(axes[0], specs):
        baseline = [v/unit for v in scenario['baseline_' + key]]
        ax.plot(range(len(baseline)), baseline, color='#0054A6', marker='o', lw=2.3, label='ML 기준 전망')
        if scenario.get('has_target'):
            ax.plot(range(len(baseline)), [v/unit for v in scenario['target_' + key]], color='#FF7400', marker='o', lw=2.3, label='목표 KPI')
        ax.set_title(label, loc='left', fontsize=11, pad=15)
        ax.set_xticks(range(len(baseline)), scenario['categories'], fontsize=8)
        ax.tick_params(axis='y', labelsize=8)
        ax.grid(axis='y', color='#DDE6EC'); ax.set_axisbelow(True)
        ax.margins(x=.14, y=.25)
        for side in ('top', 'right', 'left'): ax.spines[side].set_visible(False)
        ax.spines['bottom'].set_color('#DDE6EC')
        ax.tick_params(axis='both', length=0, pad=8)
        ax.legend(loc='upper center', bbox_to_anchor=(.5, -.2), ncol=2, frameon=False, fontsize=8)
    fig.subplots_adjust(left=.07, right=.98, top=.83, bottom=.28, wspace=.32)
    output = BytesIO(); fig.savefig(output, format='png', facecolor='white'); plt.close(fig); output.seek(0)
    return output


def metric_rows(scenario, key):
    money = key == 'spending'
    fmt = (lambda v: f'{v/100000000:,.2f}억 원') if money else (lambda v: f'{v:,.0f}명')
    rows = [['기준월', 'ML 기준 전망', '목표 KPI', '기준 대비 증가']]
    base = scenario['baseline_' + key]
    targets = scenario.get('target_' + key) if scenario.get('has_target') else None
    for month, baseline, target in zip(scenario['categories'], base, targets or [None]*len(base)):
        diff = target-baseline if target is not None else None
        change = f'↑ {diff/baseline*100:.2f}%\n+{fmt(diff)}' if diff is not None and baseline > 0 else '—'
        rows.append([month, fmt(baseline), fmt(target) if target is not None else '—', change])
    if targets:
        b, v = sum(base), sum(targets)
        rows.append(['월별 합계', fmt(b), fmt(v), f'↑ {(v/b-1)*100:.2f}%\n+{fmt(v-b)}' if b else '—'])
    return rows


def target_basis_page(doc, report, b):
    section_start = len(doc.paragraphs)
    heading(doc, '02  목표 KPI 산출근거', True)
    subheading(doc, '목표율을 정한 이유')
    reason = b['target_explanation'].split('방문 수가 0보다 큰 월은')[0].strip()
    paragraph(doc, reason, size=10).paragraph_format.space_after = Pt(10)
    scenario = report.get('execution_scenario') or {}
    paragraph(doc, f"최종월 계획 목표  방문 +{scenario.get('visitor_target_pct', 0):.2f}%  ·  소비 +{scenario.get('spending_target_pct', 0):.2f}%",
              size=15, bold=True, color=BLUE)
    subheading(doc, '3개월 동안 추가로 달성할 규모')
    rows = [['지표', '추가 목표', '월별 합계 증가율']]
    for item in b['totals']:
        value = (f"{item['additional']/1e8:,.2f}억 원" if item['key'] == 'spending'
                 else f"{item['additional']:,.0f}명")
        rows.append([item['label'], '↑ '+value, f"+{item['pct']:.2f}%" if item['pct'] is not None else '—'])
    table(doc, rows, [42, 70, 58])
    plan=(report.get('target_proposal_basis') or {}).get('capacity_plan') or {}
    if plan.get('central'):
        subheading(doc, '운영 규모와 계획 범위')
        paragraph(doc, plan['capacity_formula'], size=11, bold=True, color=BLUE)
        paragraph(doc, plan.get('schedule_note', ''), size=9, color=GRAY)
        paragraph(doc, plan['scale_rule']+' · 확보 시설 수가 아닌 운영 제안입니다.', size=9, color=GRAY)
        paragraph(doc, plan.get('participation_basis',''), size=9)
        participation=next((r for r in plan.get('rate_basis',[]) if r.get('metric')=='utilization_rate'),{})
        if participation.get('source_url'):source_link(doc, participation['source_url'], '참여 목표의 사례 근거')
        table(doc, [['시나리오', '예상 참여', '추가 방문', '추가 소비'], *[
            [s['label'], f"{s['participants']:,}명", f"{s['additional_visitors']:,}명",
             f"{s['additional_spending_krw']/10000:,.0f}만 원"] for s in plan['scenarios']]], [30,44,44,52])
        paragraph(doc, plan['spending_formula'], size=9)
        # Keep the related disclosures in one paragraph so a final sentence
        # cannot become a near-empty continuation page. All wording is kept.
        paragraph(doc, plan['disclosure'] + ' '
                  + '산출한 추가 방문·소비 규모는 각 달의 ML 전망과 단계 운영 비중에 따라 배분합니다.',
                  size=9, color=GRAY)
        # Reduce spacing only in this dense section, never its text or type.
        for p in doc.paragraphs[section_start + 1:]:
            if p.paragraph_format.space_after and p.paragraph_format.space_after > Pt(4):
                p.paragraph_format.space_after = Pt(4)
            if p.paragraph_format.space_before and p.paragraph_format.space_before > Pt(9):
                p.paragraph_format.space_before = Pt(9)
        return
    subheading(doc, '목표가 월별 수치로 이어지는 과정')
    table(doc, [['단계', '계산 방법'],
                ['01  월별 전망', '과거 월별 지표와 계절 흐름을 저장 모델에 입력해 사업기간의 전망을 계산합니다.'],
                ['02  목표 단계 적용', '첫 달부터 마지막 달까지 목표율을 단계적으로 높입니다.\n월별 목표 = 월별 전망 × (1 + 최종월 목표율 × 경과월 / 전체월)'],
                ['03  추가 규모 합산', '각 달의 목표에서 전망을 뺀 뒤 합산합니다.\n합계 증가율 = (월별 목표 합계 ÷ 월별 전망 합계 − 1) × 100']], [42, 128])
    paragraph(doc, '목표는 조정 가능한 계획 가정입니다. 사례의 사업 효과를 예측한 값이 아니며, 방문 합계에는 월간 중복이 포함됩니다.', size=9, color=GRAY)


def case_result_page(doc, report):
    result = report_outcome(report)
    if not result:
        return
    heading(doc, '03  사례 실적', True)
    subheading(doc, result['title'])
    paragraph(doc, result['scope'], size=9, color=GRAY)
    paragraph(doc, result['action'], size=12, bold=True, color=BLUE)
    photo = result.get('photo')
    if photo and photo['path'].exists():
        from PIL import Image
        with Image.open(photo['path']) as im:
            ratio = min(170/im.width, 60/im.height)
            doc.add_picture(str(photo['path']), width=Mm(im.width*ratio), height=Mm(im.height*ratio))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph(doc, photo['caption']+' · '+photo['credit'], size=8, color=GRAY)
    before, after = result['before'], result['after']
    table(doc, [['비교 기간', result['metric']],
                [result['before_label'], f"약 {before:,.0f}명"],
                [result['after_label'], f"약 {after:,.0f}명"]], [100, 70])
    paragraph(doc, f"{'↑' if after >= before else '↓'} 약 {abs(after-before):,.0f}명  ·  {(after/before-1)*100:+.2f}%", size=22, bold=True, color='E76700')
    paragraph(doc, result['before_note']+' / '+result['after_note'], size=9, color=GRAY)
    paragraph(doc, result['interpretation'], size=9, color=GRAY)
    if result.get('parent_page'):
        paragraph(doc, f"선정 문서와의 연결: 서울 문화재야행 예산 확정자료 {result['parent_page']}쪽에 중구 정동야행이 포함되어 있습니다.", size=9)
    subheading(doc, '공식 발표와 사진 출처')
    for label, url in result['references']:
        source_link(doc, url, label)


def case_card(doc, report, source, index, primary, *, add_gap=True):
    t = doc.add_table(rows=1, cols=2)
    t.autofit = False
    for col, width in zip(t.columns, (116, 54)): col.width = Mm(width)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    left, right = t.rows[0].cells
    for cell, width in zip((left, right), (116, 54)):
        cell.width = Mm(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        props = cell._tc.get_or_add_tcPr()
        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'F0F5F8'); props.append(shade)
        margins = OxmlElement('w:tcMar')
        for side in ('top', 'bottom', 'left', 'right'):
            el = OxmlElement('w:'+side); el.set(qn('w:w'), '110'); el.set(qn('w:type'), 'dxa'); margins.append(el)
        props.append(margins)
    label = '적용 사례' if source in primary else '후보 사례'
    paragraph(left, f'{index+1:02d}  {label}', size=9, bold=True, color=BLUE)
    paragraph(left, f"{source.get('case_region') or ''}  {source.get('intervention') or source.get('title') or ''}", size=11, bold=True)
    from .proposal_layout_v10 import case_card_summary
    paragraph(left, prose(case_card_summary(source)), size=9)
    source_link(left, source.get('source_url'), '사업 원문 보기')
    result = report_outcome(report) if source in primary else None
    info = (result or {}).get('photo') or case_image(source)
    if info:
        from PIL import Image
        with Image.open(info['path']) as im:
            ratio = min(46/im.width, 31/im.height)
            pic = paragraph(right)
            pic.add_run().add_picture(str(info['path']), width=Mm(im.width*ratio), height=Mm(im.height*ratio))
            pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph(right, info['caption']+' · '+info['credit'], size=7, color=GRAY)
        source_link(right, info['page_url'], '사진 출처 보기', size=7)
    for cell in (left, right):
        cell._tc.remove(cell.paragraphs[0]._p)
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(5)
    if add_gap:
        paragraph(doc, '', size=2).paragraph_format.space_after = Pt(4)


def pipeline_page(doc, report, b):
    heading(doc, '07  기획서 생성 과정', True)
    paragraph(doc, '지역 데이터와 공식 사례를 연결하고, 5개 에이전트가 조사·비교·작성·검수를 나누어 수행합니다.', size=11)
    trace = report.get('agent_trace') or []
    def provider(agent, default):
        matches = [r.get('provider') for r in trace if r.get('agent') == agent and r.get('status') == 'completed' and r.get('provider')]
        return {'openai':'OpenAI', 'qwen':'Qwen', 'gemma':'Gemma', 'ollama':'로컬 LLM'}.get(matches[-1] if matches else '', default)
    table(doc, [['흐름', '입력과 수행 내용'],
                ['기획서 생성 클릭', '선택 지역·사업기간·사용자 조건을 전달합니다.'],
                ['① 지역 근거 정리', 'Evidence Agent가 SQL 관측값·공식 관광자료·7개 ML 전망을 정리합니다.'],
                ['② 타지역 사례 조사\n'+provider('case_scout','저장 근거'), 'Case Scout가 공식 사업 내용·성과·출처를 확보합니다. ①과 ②는 병렬로 진행하며, 이미 확보한 공식 근거도 재사용합니다.'],
                ['③ 후보 비교\n'+provider('transferability','비교 모델'), 'Transferability Agent가 지역 지표와 사례의 운영 방식을 비교하여 지역에 맞는 후보를 선택합니다.'],
                ['④ 본문 작성\n'+provider('planner','작성 모델'), 'Planner Agent가 선정 후보·수치·공식 근거로 사업 소개와 실행 계획을 작성합니다.'],
                ['⑤ 품질 검수\n'+provider('reviewer','검수 모델'), 'Reviewer Agent와 코드 검사가 수치·출처·기간·본문 연결을 확인합니다. 필요한 부분은 보완 후 다시 검수합니다.'],
                ['저장과 문서 출력', '저장 보고서를 출력하며 LLM 재호출·모델 재학습은 하지 않습니다. 전국 비교 캐시가 없으면 저장 ML로 비교값을 계산합니다.']], [47, 123])
    subheading(doc, '7개 머신러닝 결과의 사용처')
    table(doc, [['입력', '결과가 사용되는 곳'],
                ['방문자 수 · 관광소비액 전망', '전망 그래프 → 월별 목표 → 추가 규모 → 견적 운영량 가정'],
                ['체류시간 · 숙박일수 · 숙박방문 비율\n내비게이션 검색 · 숙박 검색 전망', '관광 흐름 진단 → 공식 사례 조사 방향 → 후보별 지역 적용 검토']], [75, 95])
    paragraph(doc, f"저장된 공식 사례 {b['case_count']}건 → 설계 후보 {b['candidate_count']}개 → 참고 사례 {b['displayed_count']}건 표시", size=10, bold=True, color=BLUE)
    paragraph(doc, '사례 선정은 7개 예측값으로 학습한 지역 추천 모델이 아니라, 지역 지표와 공식 문서의 운영 방식을 읽고 후보를 비교하는 과정입니다. 목표와 견적은 별도의 계획 가정과 Python 산식으로 구성합니다.', size=9, color=GRAY)


def create_strategy_proposal_document(report):
    report = prepare_idea_report(report)
    strategy = (report.get('strategies') or [{}])[0]
    b = calculation_basis(report)
    scenario, demo = _scenario_for_display(report)
    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = False
    sec = doc.sections[0]
    sec.page_width = Mm(210); sec.page_height = Mm(297)
    sec.top_margin = Mm(19); sec.bottom_margin = Mm(18)
    sec.left_margin = sec.right_margin = Mm(20)
    sec.header_distance = sec.footer_distance = Mm(9)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = doc.styles[name]; style.font.name = '맑은 고딕'; style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '맑은 고딕')
        for border in list(style._element.xpath('.//w:pBdr')):
            border.getparent().remove(border)
    doc.styles['Normal'].font.size = Pt(10)
    doc.styles['Title'].paragraph_format.line_spacing = 1.05
    doc.styles['Heading 1'].paragraph_format.keep_with_next = True
    doc.styles['Heading 2'].paragraph_format.keep_with_next = True
    header = sec.header.paragraphs[0]
    header.text = f"{report.get('region_name', '')}  |  관광 전략기획안"
    header.runs[0].font.size = Pt(8); header.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('OLIGO-K   ·   ').font.size = Pt(8)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)

    paragraph(doc, report.get('region_name', ''), size=11, bold=True, color=BLUE)
    title = paragraph(doc, prose(strategy.get('title') or '관광 전략기획안'), size=23, bold=True, color='000000', style='Title')
    title.paragraph_format.line_spacing = 1.08
    title.paragraph_format.space_before = Pt(12)
    title.paragraph_format.space_after = Pt(16)
    paragraph(doc, '관광 전략기획안', size=13, color=GRAY)
    paragraph(doc, '사업기간  '+(' ~ '.join([b['months'][0], b['months'][-1]]) if b['months'] else str(strategy.get('timeframe') or '')), bold=True)
    subheading(doc, '핵심 제안')
    paragraph(doc, prose(strategy.get('summary') or report.get('summary')))
    subheading(doc, '사업의 목표')
    rows = [['지표', '3개월 목표 합계', '기준 전망 대비 추가 규모']]
    for r in b['totals']:
        fmt = (lambda v: f'{v/1e8:,.2f}억 원') if r['key']=='spending' else (lambda v: f'{v:,.0f}명')
        rows.append([r['label'], fmt(r['target']), f"↑ {r['pct']:.2f}% · +{fmt(r['additional'])}" if r['pct'] is not None else fmt(r['additional'])])
    if len(rows)>1: table(doc, rows, [38, 60, 72])
    if demo: paragraph(doc, '예시 데이터로 만든 시연용 목표입니다.', color=GRAY)
    paragraph(doc, '목표는 월별 ML 기준 전망에 계획 증가율을 적용한 값입니다. 방문 합계는 월간 순 방문 수의 합계이며, 3개월 동안 중복을 제거한 인원이 아닙니다.', size=9, color=GRAY)
    subheading(doc, '제안의 방향')
    paragraph(doc, prose(strategy.get('problem_to_solve') or strategy.get('problem')))
    paragraph(doc, prose(strategy.get('solution')))
    subheading(doc, '문서 구성')
    paragraph(doc, '01 사업 목표  ·  02 목표 KPI 산출근거  ·  03 참고 사례와 실적\n04 지역 적용  ·  05 실행 가이드  ·  06 견적  ·  07 생성 과정  ·  출처', size=9, color=GRAY)

    heading(doc, '01  사업 목표', True)
    if scenario:
        for key, name in [('visitors', '방문자 수'), ('spending', '관광소비액')]:
            if key == 'spending': heading(doc, '01  사업 목표 관광소비액', True)
            subheading(doc, name)
            doc.add_picture(forecast_chart(scenario, key), width=Mm(170))
            table(doc, metric_rows(scenario, key), [25, 47, 47, 51])
            subheading(doc, '수치 해석과 계산 기준')
            paragraph(doc, '파란선은 한국관광 데이터랩의 과거 월별 방문·소비 기록과 계절 흐름을 바탕으로 계산한 전망입니다. 주황선은 공식 사례를 참고해 정한 계획 목표입니다.', size=9)
            formula = ('준비월 추가 목표는 0명·0원입니다. 운영월 추가 목표 = 월별 전망 × 최종월 목표율 × 경과 운영월/전체 운영월.'
                       if scenario.get('operating_months') else
                       '월별 목표 = 월별 기준 전망 × (1 + 최종월 목표율 × 경과월/전체월).')
            paragraph(doc, formula + ' 합계 증가율은 월별 목표 합계 ÷ 기준 전망 합계 − 1로 계산합니다.', size=9)
            paragraph(doc, '소비액 표시는 억 원 단위 소수 둘째 자리로 반올림합니다. 합계와 증가율은 반올림 전 값으로 계산합니다. 방문 합계는 3개월 고유 인원이 아닌 월간 순 방문 수의 합계입니다.', size=8, color=GRAY)
    else:
        paragraph(doc, '해당 사업기간에 저장된 전망 수치가 없습니다. 관측값을 예측 그래프로 대체하지 않습니다.')

    heading(doc, '03  지역별 참고 사례', True)
    paragraph(doc, '공식 문서에서 사업의 실제 운영 방식이 연결되는 사례를 참고합니다. 아래 구분은 선정 사업의 근거와 추가 참고 사례의 역할을 나타냅니다.', size=9, color=GRAY)
    cases, primary = report_cases(report)
    for i, source in enumerate(cases):
        case_card(doc, report, source, i, primary, add_gap=i < len(cases) - 1)

    case_result_page(doc, report)
    heading(doc, '04  사례 선정과 지역 적용', True)
    source = primary[0] if primary else {}
    subheading(doc, '사업의 이용 흐름')
    paragraph(doc, prose(source.get('operating_model')))
    paragraph(doc, '결제·혜택·재이용 또는 예약·체류 등 실제 이용 흐름이 같은 사례를 연결합니다.', size=9, color=GRAY)
    subheading(doc, '선택 지역의 소비와 체류')
    facts = [(f.get('metric'), f.get('value')) for f in report.get('observed_findings') or [] if f.get('metric') and f.get('value')]
    if facts: table(doc, [['진단 지표', '관측값·전망']]+facts[:3], [70, 100])
    paragraph(doc, '지역 소비·체류 지표를 보고 참여 업종과 이용 대상을 정합니다.')
    subheading(doc, '운영 규모')
    paragraph(doc, b['estimate'].get('scale_basis') or '입력 예산 총액 안에서 참여량과 운영 인력을 배분합니다.')
    subheading(doc, '사례에서 우리 지역의 제안으로')
    paragraph(doc, prose(strategy.get('solution')))
    if strategy.get('comparison_analysis'):
        subheading(doc, '선정 판단')
        paragraph(doc, prose(strategy['comparison_analysis']), size=9)
    paragraph(doc, '연결 근거  '+str(source.get('title') or ''), size=8, color=GRAY)

    heading(doc, '05  4단계 실행 가이드 예시안', True)
    groups = execution_groups(report)
    steps = [['단계와 담당', '핵심 업무와 실행 방법']]
    for i, (lead, body) in enumerate(execution_copy(report)):
        steps.append([f'{i+1:02d}  {LABELS[i]}\n'+period(groups[i])+'\n'+ROLES[i], lead+'\n\n'+body])
    table(doc, steps, [45, 125])
    paragraph(doc, '일정은 저장 기획안, 업무 설명은 사업 유형별 운영 안내 함수를 사용한 실행 예시입니다. 담당 역할과 운영 조건은 협의해 조정할 수 있습니다.', size=9, color=GRAY)

    heading(doc, '06  견적 예시안', True)
    e = b['estimate']
    paragraph(doc, f"총 {e.get('total_krw', 0):,}원", size=24, bold=True, color=BLUE)
    paragraph(doc, '예상 견적이며 실제 액수와 다를 수 있습니다.', color=GRAY)
    table(doc, [['항목', '금액', '산출근거']]+[[v['name'], f"{v['amount']:,}원", v['basis']] for v in e.get('items') or []]+[['합계', f"{e.get('total_krw', 0):,}원", '항목별 금액 합계']], [36, 36, 98], green_last=True)
    subheading(doc, '규모를 정한 기준')
    paragraph(doc, e.get('scale_basis') or '사용자 지정 예산 총액을 항목별로 배분한 계획 가정입니다.')
    if e.get('scenario_note'):
        paragraph(doc, e['scenario_note'], size=10)
        paragraph(doc, f"100% 참여 시 참고예산: {e['full_participation_budget_krw']:,}원 (같은 계획 결제액·단가 기준)", size=10, bold=True, color=BLUE)
    paragraph(doc, '직접비 = 참여 혜택 + 운영·정산 + 시스템 + 홍보 + 평가. 예비비 = 직접비 × 10%. 단가와 처리량은 기획 가정입니다.', size=9)
    for note in e.get('assumptions') or []: paragraph(doc, note, size=8, color=GRAY)

    target_basis_page(doc, report, b)
    pipeline_page(doc, report, b)

    heading(doc, '참고 자료와 출처', True)
    paragraph(doc, '저장 보고서의 전체 근거 목록입니다. 원자료·ML 전망·공식 사례를 구분하고 원문 주소와 문서 위치를 유지합니다.', size=9, color=GRAY)
    for i, source in enumerate(complete_source_records(report), 1):
        p = paragraph(doc, f"{i:02d}  {source_display_name(source)}", size=9, bold=True)
        p.paragraph_format.keep_with_next = True
        details = [str(source.get('source_type') or ''), str(source.get('source_id') or '')]
        if source.get('page_references'): details.append('페이지 '+str(source['page_references']))
        if source.get('chunk_ids'): details.append('문서 구간 '+', '.join(source['chunk_ids']))
        p.paragraph_format.space_after = Pt(2)
        p = paragraph(doc, ' · '.join(filter(None, details)), size=8, color=GRAY)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = bool(source.get('source_url'))
        if source.get('source_url'):
            source_link(doc, source['source_url'])
    # Keep the reasons beside the goals; page count follows content length.
    body = doc._element.body
    children = list(body)
    start = next(i for i,e in enumerate(children) if e.tag == qn('w:p') and ''.join(e.itertext()).startswith('02  목표 KPI'))
    end = next(i for i,e in enumerate(children) if i > start and e.tag == qn('w:p') and ''.join(e.itertext()).startswith('07  기획서'))
    anchor = next(e for e in children if e.tag == qn('w:p') and ''.join(e.itertext()).startswith('03  지역별'))
    for element in children[start:end]: anchor.addprevious(element)
    # Compact supporting sources while preserving every full title and URL.
    source_mode = False
    for p in doc.paragraphs:
        if p.text == '참고 자료와 출처': source_mode = True; continue
        if source_mode:
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.space_after = Pt(4)
    out = BytesIO(); doc.save(out); out.seek(0)
    return out

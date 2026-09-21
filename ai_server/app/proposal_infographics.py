"""Editable calculation diagrams drawn from the report, never invented values."""
from .proposal_layout_v7 import text, rect, header, BLUE, CYAN, ORANGE, SLATE, PALE, WHITE
from .proposal_calculation_basis import calculation_basis


def arrow(slide,name,x,y,w=72):
    text(slide,name,'→',x,y,w,65,42,CYAN,True)


def metric(slide,name,label,value,sub,x,y,w=380):
    text(slide,name+'-label',label,x,y,w,38,24,BLUE,True)
    text(slide,name+'-value',value,x,y+48,w,60,38,ORANGE,True)
    text(slide,name+'-sub',sub,x,y+117,w,76,21,SLATE)


def icon(slide, name, kind, x, y, size=72, color='blue'):
    from pathlib import Path
    from .proposal_layout_v7 import EMU
    path=Path(__file__).resolve().parents[1]/'assets'/'proposal_icons'/f'{kind}-{color}.png'
    shape=slide.shapes.add_picture(str(path),int(x*EMU),int(y*EMU),int(size*EMU),int(size*EMU))
    shape.name=name
    return shape


def route(slide, name, points, color=CYAN, head=True):
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.dml.color import RGBColor
    from pptx.util import Pt
    from .proposal_layout_v7 import EMU
    for i,((x1,y1),(x2,y2)) in enumerate(zip(points,points[1:])):
        sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,int(x1*EMU),int(y1*EMU),int(x2*EMU),int(y2*EMU))
        sh.name=f'{name}-{i}';sh.line.color.rgb=RGBColor.from_string(color);sh.line.width=Pt(1.7)
        if head and i==len(points)-2:
            end=OxmlElement('a:tailEnd');end.set('type','triangle');sh.line._get_or_add_ln().append(end)


def centered_copy(slide,name,value,x,y,w,h,size=24,color=BLUE,bold=False):
    from pptx.enum.text import PP_ALIGN
    sh=text(slide,name,value,x,y,w,h,size,color,bold)
    for p in sh.text_frame.paragraphs:p.alignment=PP_ALIGN.CENTER
    return sh


def calculations(slide, report):
    b=calculation_basis(report); totals={r['key']:r for r in b['totals']}
    scenario=report.get('execution_scenario') or {}; basis=report.get('target_proposal_basis') or {}
    visitor=scenario.get('visitor_target_pct',0);spend=scenario.get('spending_target_pct',0)
    header(slide,'목표 KPI 산출근거')
    text(slide,'kpi-intro','목표율을 정한 이유와 3개월 동안 추가로 달성할 규모',106,211,1388,47,27,SLATE)
    icon(slide,'kpi-reason-icon','target',126,304,85)
    text(slide,'kpi-reason-title','목표를 정한 이유',236,310,568,56,33,BLUE,True)
    default_v=basis.get('visitor_target_pct',5);default_s=basis.get('spending_target_pct',5)
    plan=basis.get('capacity_plan') or {}
    if plan.get('central') and basis.get('target_mode') == 'automatic':
        c=plan['central']
        region=report.get('region_name') or '선택 지역'
        period=plan.get('operating_period') or plan.get('forecast_period','')
        period=' ~ '.join(f'{p[:4]}.{p[4:]}' if len(p)==6 and p.isdigit() else p
                          for p in period.split('~'))
        reason=(f"{region}의 {plan['months']}개월 운영안\n"
                f"운영 산정 기간 {period}\n"
                f"편성 정원 {plan['funded_capacity']:,}명분 × 이용률 {c['utilization_pct']:g}%\n"
                f"예상 참여 {c['participants']:,}명 중 추가 방문 {c['additional_visitor_share_pct']:g}% 가정")
        decision=f"추가 방문 목표 {c['additional_visitors']:,}명"
    elif basis.get('target_mode') == 'user':
        reason='현재 목표는 사용자가 지정한 값입니다.\n운영 규모 시나리오는 다음 장의 참고안입니다.'
        decision=f'방문 +{visitor:.2f}% · 소비 +{spend:.2f}% 적용'
    elif basis.get('kind') == 'operating_capacity':
        reason='사업기간 전망과 운영 조건을 연결해\n참여 규모를 검토합니다.'
        decision='운영 조건에 따른 목표 제안'
    elif visitor!=default_v or spend!=default_s:
        reason='공식 사례를 참고한 초기 제안에서\n현재 저장된 목표율로 조정했습니다.'
        decision=f'방문 +{visitor:g}% · 소비 +{spend:g}% 적용'
    elif basis.get('observed_visitors') and basis.get('reported_growth_pct'):
        growth=basis['reported_growth_pct']
        reason=f'강진의 발표 증가율 {growth:g}%에\n단계 운영을 고려한 계획 계수를 적용했습니다.'
        decision=f'{growth:g}% × 80% = {visitor:g}%' if visitor==20 and growth==25 else f'방문 +{visitor:g}% · 소비 +{spend:g}% 채택'
    else:
        reason='사례의 운영 방식은 참고하되, 성장률을\n그대로 옮기지 않고 초기 목표 5%를 채택했습니다.'
        decision=f'조정 가능한 최종월 목표 +{visitor:g}%'
    automatic_capacity=bool(plan.get('central') and basis.get('target_mode') == 'automatic')
    text(slide,'kpi-reason',reason,126,398 if automatic_capacity else 408,698,
         138 if automatic_capacity else 113,26 if automatic_capacity else 31,SLATE)
    rect(slide,'kpi-decision-band',126,550,671,110,PALE)
    text(slide,'kpi-decision',decision,148,573,627,63,38,BLUE,True)
    step=('추가 방문·소비 규모를 먼저 계산하고\n월별 전망과 단계 운영 비중에 따라 배분합니다.' if plan.get('central') else
          '첫 달부터 마지막 달까지 목표율을\n단계적으로 높여 3개월 추가 목표를 합산합니다.')
    if automatic_capacity:
        step=(f"운영 정원 = {plan['sites']}개 거점 × 월 {plan['days_per_month']}일 × {plan['months']}개월\n"
              f"× 하루 {plan['sessions_per_day']}회 × 회당 {plan['capacity_per_session']}명 = {plan['capacity']:,}명분\n"
              '운영량은 기획 가정이며 참여 비율의 근거는 다음 장에 표시합니다.')
    text(slide,'kpi-step-explanation',step,126,695 if automatic_capacity else 706,671,
         103 if automatic_capacity else 84,23 if automatic_capacity else 25,SLATE)
    route(slide,'kpi-to-visitors',[(815,560),(864,560),(864,409),(913,409)])
    route(slide,'kpi-to-spending',[(864,560),(864,670),(913,670)])
    for i,(key,label,symbol) in enumerate([('visitors','추가 방문 목표','users'),('spending','추가 소비 목표','coins')]):
        row=totals.get(key); y=292+i*261
        rect(slide,f'kpi-impact-panel-{i}',913,y,581,233,WHITE)
        icon(slide,f'kpi-impact-icon-{i}',symbol,939,y+24,55)
        text(slide,f'kpi-impact-title-{i}',label,1011,y+25,453,49,29,BLUE,True)
        value=(f"{row['additional']:,.0f}명" if key=='visitors' else f"{row['additional']/1e8:,.2f}억 원") if row else '저장 목표 확인'
        text(slide,f'kpi-impact-value-{i}','↑ '+value,938,y+95,535,78,49,ORANGE,True)
        pct=f"{row['pct']:+.2f}%" if row and row['pct'] is not None else ''
        text(slide,f'kpi-impact-period-{i}',f'3개월 월별 합계 · {pct}',940,y+181,530,35,22,SLATE)
    text(slide,'kpi-footnote','목표는 계획 가정입니다. 사례의 사업 효과를 예측한 값이 아니며, 방문 합계에는 월간 중복이 포함됩니다.',106,850,1388,39,20,SLATE)


def ml_to_plan(slide, report):
    header(slide,'기획서 생성 파이프라인')
    text(slide,'pipeline-intro','5개 전문 에이전트가 조사·비교·작성·검수를 나누어 수행합니다',106,211,1388,47,27,SLATE)
    trace=report.get('agent_trace') or []
    def provider(agent,fallback):
        rows=[r for r in trace if r.get('agent')==agent and r.get('status')=='completed' and r.get('provider')]
        raw=rows[-1]['provider'] if rows else ''
        return {'openai':'OpenAI','qwen':'AI','gemma':'AI','ollama':'로컬 LLM'}.get(raw,fallback)
    case_provider=provider('case_scout','저장 근거')
    # Lines are behind the icons and labels; orthogonal forks show parallel work.
    route(slide,'start-to-fork',[(258,454),(283,454)],head=False)
    route(slide,'fork-evidence',[(283,454),(283,354),(310,354)])
    route(slide,'fork-case',[(283,454),(283,554),(310,554)])
    route(slide,'join-evidence',[(600,354),(622,354),(622,454)],head=False)
    route(slide,'join-case',[(600,554),(622,554),(622,454)],head=False)
    route(slide,'join-compare',[(622,454),(651,454)])
    for i,(a,z) in enumerate([(832,874),(1054,1096),(1276,1340)]):route(slide,f'pipeline-next-{i}',[(a,454),(z,454)])
    rect(slide,'research-background',310,278,290,363,PALE)
    for i,(name,kind,title,subtitle) in enumerate([
        ('evidence','map-pinned','① 지역 근거 정리','SQL · 공식 관광자료'),
        ('case','search-check','② 타지역 사례 조사',case_provider+' · 공식 사업·성과')]):
        y=305+i*198
        icon(slide,f'pipeline-{name}-icon',kind,330,y+8,57)
        text(slide,f'pipeline-{name}-title',title,399,y,189,67,24,BLUE,True)
        centered_copy(slide,f'pipeline-{name}-subtitle',subtitle,324,y+82,261,51,20,SLATE)
    icon(slide,'pipeline-click-icon','mouse-pointer-click',142,365,86)
    centered_copy(slide,'pipeline-click','기획서 생성\n클릭',106,483,150,91,27,BLUE,True)
    stages=[(652,'git-compare-arrows','③ 후보 비교',provider('transferability','비교 모델'),'지역에 맞는\n사업 선택'),
            (874,'notebook-pen','④ 본문 작성',provider('planner','작성 모델'),'선정 근거로\n실행 계획 작성'),
            (1096,'shield-check','⑤ 품질 검수',provider('reviewer','검수 모델'),'수치·출처와\n본문 연결 확인')]
    for x,kind,title,model,body in stages:
        centered_copy(slide,f'pipeline-model-{x}',model,x+10,308,158,48,27,BLUE,True)
        icon(slide,f'pipeline-icon-{x}',kind,x+49,389,80)
        centered_copy(slide,f'pipeline-title-{x}',title,x+9,496,160,48,25,BLUE,True)
        centered_copy(slide,f'pipeline-body-{x}',body,x+9,554,160,77,22,SLATE)
    icon(slide,'pipeline-output-icon','files',1372,389,76)
    centered_copy(slide,'pipeline-output-title','저장·출력',1340,496,154,48,25,BLUE,True)
    centered_copy(slide,'pipeline-output-body','웹\nPPT · Word',1340,554,154,77,22,SLATE)
    # A shared input rail makes the data contribution visible without another
    # wall of model descriptions. It feeds investigation and candidate comparison.
    rect(slide,'pipeline-data-band',106,710,694,101,WHITE)
    icon(slide,'pipeline-data-icon','database',128,731,58)
    text(slide,'pipeline-data-label','지역 관측값 + 7개 ML 전망',207,723,566,40,26,BLUE,True)
    text(slide,'pipeline-data-use','조사 방향과 후보 비교에 함께 전달',207,765,566,33,21,SLATE)
    route(slide,'data-to-research',[(455,709),(455,658)])
    route(slide,'data-to-compare',[(741,709),(741,647)])
    text(slide,'pipeline-revision','필요한 부분은 보완 작성 후 다시 검수',857,727,637,39,23,SLATE)
    route(slide,'review-return',[(1185,642),(1185,688),(963,688),(963,642)],color='8CA4BF')
    text(slide,'pipeline-footnote','OpenAI는 공식 웹에서 사업 사례와 출처를 조사합니다. 이미 확보한 공식 근거는 재사용합니다.',106,850,1388,39,19,SLATE)


def planning_rationale_page(slide, report):
    """A saved candidate's facts, hypothesis and common comparison criteria."""
    from .planning_rationale import report_rationale
    r = report_rationale(report)
    if not r:
        raise ValueError('선정 후보에 연결된 판단 기록이 필요합니다.')
    header(slide, '선정 근거 ③ 사실에서 적용 가설까지')
    reasoning = r['reasoning']
    text(slide, 'rationale-facts-label', '01  연결한 관측·전망', 106, 210, 652, 42, 28, BLUE, True)
    text(slide, 'rationale-hypothesis-label', '02  우리 지역에 적용할 가설', 817, 210, 677, 42, 28, BLUE, True)
    names = {'visitors': '방문자 수', 'spending_krw': '관광소비액', 'stay_minutes': '체류시간',
             'lodging_nights': '숙박일수', 'lodging_rate_pct': '숙박방문 비율',
             'navigation_searches': '내비게이션 검색', 'lodging_searches': '숙박 검색'}
    for i, fact in enumerate(r['facts'].values()):
        row = fact['record']
        if fact['kind'] == 'observed':
            value = '관측 · ' + str(row.get('summary') or row.get('title') or '')
        else:
            number = row.get('forecast_value')
            unit = str(row.get('unit') or '')
            formatted = f'{number:,.2f}' if isinstance(number, (int, float)) else str(number or '')
            if isinstance(number, (int, float)) and unit == '원' and abs(number) >= 100_000_000:
                formatted, unit = f'{number/100_000_000:,.2f}', '억 원'
            elif isinstance(number, (int, float)) and unit in ('명', '건'):
                formatted = f'{number:,.0f}'
            value = f"ML 전망 · {names.get(row.get('metric'), row.get('metric'))}\n{row.get('period')} · {formatted}{unit}"
            if row.get('aggregation') == '3_month_sum':
                value += ' · 3개월 합계'
            elif row.get('aggregation') == '3_month_average':
                value += ' · 3개월 평균'
        text(slide, f'rationale-fact-{i}', f'[{i+1}] {value}', 106, 269+i*72, 652, 68, 22, SLATE)
    text(slide, 'rationale-hypothesis', reasoning.get('application_hypothesis'), 817, 269, 677, 148, 27, BLUE, True)
    sources = {s.get('source_id'): s for s in report.get('evidence_sources') or []}
    case_names = [str(sources[sid].get('case_region') or sources[sid].get('title') or sid)
                  for sid in r['case_source_ids'] if sid in sources]
    text(slide, 'rationale-cases', '운영 참고 · ' + ' / '.join(case_names), 817, 424, 677, 67, 21, SLATE)
    rect(slide, 'rationale-mid-rule', 106, 506, 1388, 2, CYAN)
    text(slide, 'rationale-criteria-title', '03  모든 후보에 적용한 세 가지 비교 기준', 106, 521, 1388, 41, 28, BLUE, True)
    for i, (key, title) in enumerate([('visitor_action', '이용 행동'), ('operating_requirements', '운영 준비'), ('measurement', '성과 확인')]):
        x = 106+i*478
        text(slide, f'rationale-criterion-{i}', title, x, 583, 431, 35, 25, BLUE, True)
        text(slide, f'rationale-detail-{i}', reasoning.get(key), x, 626, 431, 87, 23, SLATE)
    rect(slide, 'rationale-selection-band', 106, 730, 1388, 126, PALE)
    text(slide, 'rationale-choice-label', '선택 이유', 123, 744, 171, 40, 25, BLUE, True)
    choice = text(slide, 'rationale-choice', r['selection_reason'], 306, 740, 1138, 106, 20, BLUE)
    # Do not combine pre-wrapped paragraphs and native wrapping: different font
    # metrics can turn each short trailing word into an extra line.
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    fitted_size = min(15, choice.text_frame.paragraphs[0].runs[0].font.size.pt)
    choice.text_frame.clear()
    choice.text_frame.word_wrap = True
    p = choice.text_frame.paragraphs[0]
    p.space_before = p.space_after = Pt(0)
    p.line_spacing = Pt(fitted_size * 7 / 6)
    run = p.add_run()
    run.text = r['selection_reason']
    run.font.name = 'Noto Sans KR'
    run.font.size = Pt(fitted_size)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    text(slide, 'rationale-meaning', '위 연결은 기획 판단입니다. 사업 효과는 가설이며, 실제 이용 결과로 확인합니다. 출처·비교 원문은 발표자 메모에 보존합니다.', 106, 867, 1388, 30, 18, SLATE)

"""Final user-approved adjustments after all legacy index-based layouts."""
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from .proposal_layout_v7 import EMU, text, rect, fit_text
from .proposal_theme_spicus import get,pos,center,rounded
from .proposal_infographics import icon


def remove(slide,name):
    sh=get(slide,name)
    if sh is not None:sh._element.getparent().remove(sh._element)


def toc_typography(prs):
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.name.startswith('toc-') and sh.name.endswith(('-number','-label')) and sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    para.line_spacing=Pt(25)
                    for run in para.runs:run.font.size=Pt(20)


def reference_card_backgrounds(prs):
    for slide in prs.slides:
        for n in range(3):
            box=get(slide,f'case-panel-{n}')
            if box is None:continue
            bottom=box.top+box.height
            box.top=196*EMU
            box.height=bottom-box.top
            rounded(box,'DCEBF8' if n==0 else 'FFFFFF')


def kpi_content_typography(prs):
    sizes={'kpi-reason-title':20,'kpi-reason':16,'kpi-decision':20,
           'kpi-step-explanation':15,'kpi-footnote':13.5}
    for slide in prs.slides:
        if get(slide,'kpi-reason-card') is None:continue
        for sh in slide.shapes:
            size=sizes.get(sh.name)
            if sh.name.startswith('kpi-impact-title-'):size=18
            elif sh.name.startswith('kpi-impact-value-'):size=27
            elif sh.name.startswith('kpi-impact-period-'):size=13.5
            if size is None or not sh.has_text_frame:continue
            for para in sh.text_frame.paragraphs:
                para.line_spacing=Pt(size*1.3)
                for run in para.runs:run.font.size=Pt(size)


def final_polish(prs):
    from .proposal_outlook_cards import restyle_outlook_cards
    restyle_outlook_cards(prs)
    kpi_content_typography(prs)
    toc_typography(prs)
    reference_card_backgrounds(prs)
    # Swap the two overview pages, and keep section numbering in reading order.
    overview=next((s for s in prs.slides if get(s,'overview-description') is not None),None)
    outlook=next((s for s in prs.slides if get(s,'title') is not None and '지역 관광 전망' in get(s,'title').text),None)
    if overview is not None and outlook is not None:
        ids=prs.slides._sldIdLst
        a=next(x for x in ids if x.id==overview.slide_id)
        b=next(x for x in ids if x.id==outlook.slide_id)
        ai,bi=list(ids).index(a),list(ids).index(b)
        if ai<bi:
            ids.remove(b);ids.insert(ai,b)
        fit_text(get(outlook,'title'),'1.1 지역 관광 전망',40,'22262A',True)
        fit_text(get(overview,'title'),'1.2 제안 사업 소개',40,'22262A',True)
    for i,slide in enumerate(prs.slides,1):
        for sh in slide.shapes:
            if sh.name in ('editorial-kicker','title') and sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    para.alignment=PP_ALIGN.LEFT
                    for run in para.runs:run.font.size=Pt(17.3 if sh.name=='editorial-kicker' else 30)
            if sh.name=='editorial-page':
                fit_text(sh,f'{i:02d} / {len(prs.slides):02d}',15,'66717A');center(sh)
        if get(slide,'result-flow-panel-0') is not None:
            case_flow(slide)
        if get(slide,'pipeline-click-icon') is not None:
            old=get(slide,'pipeline-click-icon');x,y,size=old.left/EMU,old.top/EMU,old.width/EMU
            remove(slide,old.name);icon(slide,'pipeline-click-icon','mouse-pointer-click',x,y,size,'red')
            remove(slide,'pipeline-output-icon')
            for j,(letter,color,label) in enumerate([('P','D35230','PowerPoint'),('W','185ABD','Word')]):
                x=1342+j*83
                doc=rect(slide,f'pipeline-export-doc-{letter}',x+8,383,62,73,color);rounded(doc,color,7000)
                rect(slide,f'pipeline-export-line-{letter}',x+23,399,32,3,'FFFFFF')
                rect(slide,f'pipeline-export-line2-{letter}',x+23,409,32,3,'FFFFFF')
                tile=rect(slide,f'pipeline-export-tile-{letter}',x,416,46,44,color);rounded(tile,color,5000)
                sh=text(slide,f'pipeline-export-letter-{letter}',letter,x+4,419,38,37,31,'FFFFFF',True);center(sh)
                sh.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE
                sh=text(slide,f'pipeline-export-label-{letter}',label,x-3,466,81,23,12,'62666C');center(sh)


def case_flow(slide):
    remove(slide,'result-scope');remove(slide,'result-operation-label')
    panel=get(slide,'result-photo-panel');pos(panel,106,250,586,365)
    photo=get(slide,'result-photo')
    y,height=250,365
    if photo is not None:
        iw,ih=photo.image.size;scale=min(586/iw,365/ih)
        w,height=iw*scale,ih*scale;y=250+(365-height)/2
        pos(photo,106+(586-w)/2,y,w,height)
        bg=get(slide,'result-photo-source-bg');cap=get(slide,'result-photo-source')
        if bg is not None and cap is not None:
            h=max(46,bg.height/EMU)
            pos(bg,photo.left/EMU+12,y+height-h-12,w-24,h)
            pos(cap,photo.left/EMU+24,y+height-h-9,w-48,h-6)
            fit_text(cap,cap.text,18,'FFFFFF');center(cap);cap.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE
    # Panoramic reference photos must not collapse the three readable flow rows.
    flow_height=max(300,height)
    y=250+(365-flow_height)/2
    row_height=(flow_height-20)/3
    for n in range(3):
        top=y+n*(row_height+10)
        pos(get(slide,f'result-flow-panel-{n}'),753,top,741,row_height)
        pos(get(slide,f'result-flow-icon-{n}'),777,top+(row_height-48)/2,48,48)
        sh=get(slide,f'result-flow-title-{n}');pos(sh,845,top+10,623,39);fit_text(sh,sh.text,26,'26323C',True)
        sh=get(slide,f'result-flow-detail-{n}');pos(sh,845,top+52,623,row_height-57);fit_text(sh,sh.text,22,'62666C')
        remove(slide,f'result-flow-arrow-{n}-0')

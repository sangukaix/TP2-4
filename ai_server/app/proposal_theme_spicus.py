"""SPICUS reference-inspired editorial theme; numeric objects remain native.

Reference: user-supplied 2024 corporate proposal, slides 3/20/25/28.
Only visual vocabulary is adapted. No SPICUS claims or branding are imported.
"""
from hashlib import sha256
from io import BytesIO
from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

from .proposal_layout_v7 import EMU, text, rect, fit_text, clear

RED = 'F50045'
DARK = '202124'
TEAL = '36A69C'
GRAY = 'F3F3F3'
PALE = 'FFF0F3'
MUTED = '62666C'
WHITE = 'FFFFFF'
CAPTION_BLUE = '0F315B'
OCEAN = '27649B'
JADE = '187C70'
VIOLET = '7556A5'
OCHRE = 'A96924'
SLATE = '495565'
ASSETS = Path(__file__).resolve().parents[1] / 'assets' / 'proposal_icons'


def pos(shape, x, y, w=None, h=None):
    shape.left, shape.top = round(x*EMU), round(y*EMU)
    if w is not None: shape.width = round(w*EMU)
    if h is not None: shape.height = round(h*EMU)


def rounded(shape, fill=WHITE, radius=10000):
    geometry = shape._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom')
    if geometry is not None:
        geometry.set('prst', 'roundRect')
        av = geometry.find('{http://schemas.openxmlformats.org/drawingml/2006/main}avLst')
        if av is None:
            av = OxmlElement('a:avLst'); geometry.append(av)
        for child in list(av): av.remove(child)
        guide = OxmlElement('a:gd'); guide.set('name','adj'); guide.set('fmla',f'val {radius}'); av.append(guide)
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    shape.line.fill.background()


def panel(slide, name, x, y, w, h, fill=WHITE):
    shape = rect(slide, name, x, y, w, h, fill)
    rounded(shape, fill)
    el=shape._element; el.getparent().remove(el); slide.shapes._spTree.insert(2,el)
    return shape


def center(shape):
    for p in shape.text_frame.paragraphs: p.alignment=PP_ALIGN.CENTER
    shape.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE


def get(slide, name):
    return next((s for s in slide.shapes if s.name==name),None)


def delete(slide, names):
    for s in list(slide.shapes):
        if s.name in names:
            s._element.getparent().remove(s._element)


def recolor_xml(element, mapping):
    for color in element.xpath('.//a:srgbClr'):
        value=color.get('val','').upper()
        if value in mapping: color.set('val',mapping[value])


def charts(slide):
    for shape in slide.shapes:
        if not shape.has_chart: continue
        chart=shape.chart
        recolor_xml(chart._chartSpace, {'004EA2':RED,'00B7C9':TEAL,'FF6B00':RED,'0054AA':TEAL})
        if shape.name.startswith('comparison-chart-'):
            for series,color in zip(chart.series,(TEAL,RED)):
                series.format.line.color.rgb=RGBColor.from_string(color)
                series.format.line.width=Pt(2.5)
                for node in series._element.xpath('.//a:srgbClr'): node.set('val',color)
        for node in chart._chartSpace.xpath('.//a:defRPr'):
            # Existing native chart labels keep their positions and values.
            for latin in node.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}latin'):
                latin.set('typeface','Noto Sans KR')


def recolor_icons(slide):
    originals={sha256(p.read_bytes()).digest():p.stem.rsplit('-',1)[0] for p in ASSETS.glob('*-blue.png')}
    for shape in list(slide.shapes):
        if not hasattr(shape,'image'): continue
        kind=originals.get(sha256(shape.image.blob).digest())
        if not kind or not (ASSETS/f'{kind}-red.png').exists(): continue
        tree=shape._element.getparent(); index=tree.index(shape._element)
        new=slide.shapes.add_picture(str(ASSETS/f'{kind}-red.png'),shape.left,shape.top,shape.width,shape.height)
        new.name=shape.name
        tree.remove(shape._element);tree.remove(new._element);tree.insert(index,new._element)


def cover(slide, photo, ending=False, region_name='', photo_source=None):
    copy={s.name:s.text for s in slide.shapes if s.has_text_frame}
    clear(slide)
    slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string(RED if not ending else DARK)
    if photo:
        picture=slide.shapes.add_picture(BytesIO(photo),940*EMU,0,660*EMU,900*EMU)
        picture.name='cover-regional-photo'
        # Preserve photographic aspect ratio through a centered crop.
        iw,ih=picture.image.size; ratio=iw/ih; box=660/900
        if ratio>box:
            crop=(1-box/ratio)/2;picture.crop_left=picture.crop_right=crop
        else:
            crop=(1-ratio/box)/2;picture.crop_top=picture.crop_bottom=crop
    rect(slide,'cover-top-rule',80,95,70,5,WHITE)
    text(slide,'cover-category',copy.get('security-label','지역 관광 전략기획안'),80,130,780,48,26,WHITE,True)
    first=copy.get('cover-title-line1',''); second=copy.get('cover-title-line2','')
    text(slide,'cover-title-line1',first,80,275,810,125,58,WHITE,True)
    text(slide,'cover-title-line2',second,80,414,810,142,48,WHITE,True)
    text(slide,'cover-date',copy.get('cover-date',''),80,610,800,44,25,WHITE)
    text(slide,'org-title',copy.get('org-title',''),80,760,780,48,28,WHITE,True)
    details=' · '.join(copy[k].replace('\n',' ') for k in ('org-owner','report-number','report-topic') if copy.get(k))
    text(slide,'cover-meta',details,80,814,790,60,16,WHITE)
    if not ending:
        # User reference: date, province ribbon, left-aligned project, period.
        region_label=region_name.strip() or copy.get('org-title','')
        delete(slide,{'cover-category','cover-title-line2','cover-meta'})
        pos(get(slide,'cover-top-rule'),80,82,170,4)
        import re
        date_match=re.search(r'\d{4}-\d{2}-\d{2}',details)
        if date_match:
            text(slide,'cover-issued-date',date_match.group().replace('-','.'),84,137,750,46,28,WHITE,True)
        def set_alpha(shape,alpha):
            color=shape._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
            if color is not None:
                opacity=OxmlElement('a:alpha');opacity.set('val',str(alpha));color.append(opacity)
        def ribbon(name,points,color,alpha=100000):
            builder=slide.shapes.build_freeform(points[0][0]*EMU,points[0][1]*EMU)
            builder.add_line_segments([(x*EMU,y*EMU) for x,y in points[1:]],close=True)
            shape=builder.convert_to_shape();shape.name=name
            shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(color)
            set_alpha(shape,alpha);shape.line.fill.background()
            return shape
        ribbon_panel=rect(slide,'cover-ribbon-panel',398,196,542,72,WHITE)
        set_alpha(ribbon_panel,50000)
        ribbon('cover-ribbon-main',[(0,196),(834,196),(678,268),(0,268)],WHITE)
        ribbon('cover-ribbon-soft',[(0,196),(931,196),(772,268),(0,268)],WHITE,23000)
        org=get(slide,'org-title');pos(org,64,196,510,72);fit_text(org,region_label,40,'000000',True)
        org.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE
        org.text_frame.margin_top=org.text_frame.margin_bottom=0
        # Keep the region label above the ribbon shapes in drawing order.
        org._element.getparent().remove(org._element);slide.shapes._spTree.append(org._element)
        project=' '.join((first,second)).strip()
        title=get(slide,'cover-title-line1');pos(title,74,313,790,215);fit_text(title,project,43,WHITE,True)
        # Rule follows actual wrapped title lines; long regional titles remain inside.
        lines=len(title.text_frame.paragraphs)
        rect(slide,'cover-title-rule',80,max(450,313+lines*58+30),170,4,WHITE)
        date=get(slide,'cover-date');pos(date,80,739,790,52);fit_text(date,copy.get('cover-date',''),27,WHITE)
    if ending:
        delete(slide,{'cover-category','cover-date'})
    if photo and photo_source:
        from .proposal_layout_v7 import wrap_words
        place=f"{region_name} · {photo_source.get('title','')}"
        caption=place
        lines=wrap_words(caption,556,17)
        height=max(42,len(lines)*22+16);y=882-height
        rect(slide,'cover-source-bg',972,y,596,height,CAPTION_BLUE)
        label=text(slide,'cover-source-label',caption,983,y+4,574,height-8,17,WHITE)
        center(label)
        link=photo_source.get('source_url') or photo_source.get('image_url')
        if link:label.click_action.hyperlink.address=link
        import json
        slide.notes_slide.notes_text_frame.text += '\n표지 사진 출처: '+json.dumps(photo_source,ensure_ascii=False)
    if not ending:
        soften_cover_background(slide)


def soften_cover_background(slide):
    """Keep the approved navy cover and place the regional photo in the right column."""
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb=RGBColor.from_string('203B4B')
    photo=get(slide,'cover-regional-photo')
    if photo is not None:
        pos(photo,940,0,660,900)
        photo.crop_left=photo.crop_right=photo.crop_top=photo.crop_bottom=0
        iw,ih=photo.image.size; ratio=iw/ih; box=660/900
        if ratio>box:photo.crop_left=photo.crop_right=(1-box/ratio)/2
        else:photo.crop_top=photo.crop_bottom=(1-ratio/box)/2

    delete(slide,{'cover-photo-blend','cover-province-ribbon','cover-ribbon-light','cover-ribbon-pink'})
    org=get(slide,'org-title')
    if org is not None:
        pos(org,64,196,510,72)
        fit_text(org,org.text,40,'000000',True)
        org.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE
        org.text_frame.margin_top=org.text_frame.margin_bottom=0
    for name in ('cover-top-rule','cover-title-rule'):
        rule=get(slide,name)
        if rule is not None:rule.fill.fore_color.rgb=RGBColor.from_string('A9D7D1')


def contents(slide):
    labels=[(get(slide,f'toc-{i}-number').text,get(slide,f'toc-{i}-label').text) for i in range(7)]
    brand=get(slide,'top-brand').text
    clear(slide);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string(GRAY)
    rail=rect(slide,'contents-rail',0,0,1600,900,'E5E7EB')
    rail.fill.gradient()
    rail.fill.gradient_angle=0
    for stop,color in zip(rail.fill.gradient_stops,('E7E9ED','FAFAFB')):
        stop.color.rgb=RGBColor.from_string(color)
    text(slide,'contents-en','CONTENTS',106,102,350,34,18,SLATE)
    text(slide,'contents-ko','목차',106,148,900,75,42,DARK,True)
    rect(slide,'contents-short-rule',112,244,58,3,RED)
    text(slide,'top-brand',brand.replace('\n',' '),106,799,1200,42,17,SLATE)
    for i,(number,label) in enumerate(labels):
        x=106 if i<4 else 865
        y=315+(i if i<4 else i-4)*101
        text(slide,f'toc-{i}-number',number,x,y,50,46,24,RED,True)
        text(slide,f'toc-{i}-label',label,x+68,y+1,565,48,24,DARK)
        rect(slide,f'contents-rule-{i}',x+6,y+69,612,1,'CDD2D8')


def growth(slide):
    title=get(slide,'title');pos(title,552,87,942,80);fit_text(title,title.text,42,DARK,True)
    for p in title.text_frame.paragraphs:p.alignment=PP_ALIGN.LEFT
    if get(slide,'purpose-period'):
        pos(get(slide,'purpose-period'),566,205,920,45)
        center(get(slide,'purpose-period'))
    pos(get(slide,'purpose-comparison'),566,253,920,40) if get(slide,'purpose-comparison') else None
    if get(slide,'purpose-comparison'):center(get(slide,'purpose-comparison'))
    for i in range(2):
        label=get(slide,f'growth-label-{i}')
        if label is None:continue
        y=312+i*246
        panel(slide,f'growth-card-{i}',551,y,950,229)
        pos(label,578,y+13,430,44);fit_text(label,label.text,26,DARK,True)
        total=get(slide,f'growth-total-{i}');pos(total,1000,y+17,474,38);fit_text(total,total.text,21,MUTED)
        chart=get(slide,f'growth-chart-{i}');pos(chart,578,y+60,896,112)
        chart.chart.plots[0].gap_width=115
        gap=get(slide,f'growth-gap-{i}');pos(gap,578,y+178,895,39);fit_text(gap,gap.text,22,RED,True)
    cap=get(slide,'hero-source-bg')
    if cap:rounded(cap,CAPTION_BLUE,5000)
    note=get(slide,'purpose-note')
    if note:pos(note,566,828,921,68)


def tables(slide):
    for shape in slide.shapes:
        if not shape.has_table:continue
        for ri,row in enumerate(shape.table.rows):
            for ci,cell in enumerate(row.cells):
                color=RED if ri==0 else (WHITE if ri%2 else 'FBEFF2')
                if shape.name=='estimate-table' and ri==0 and ci==2:color=TEAL
                cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(color)
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.color.rgb=RGBColor.from_string(WHITE if ri==0 else DARK)


def balanced_colors(slide, index):
    """Keep red for section identity; use neutral text and colors by diagram role."""
    def ink(name,color):
        shape=get(slide,name)
        if shape is not None and shape.has_text_frame:
            for node in shape._element.xpath('.//a:rPr/a:solidFill/a:srgbClr | .//a:defRPr/a:solidFill/a:srgbClr | .//a:endParaRPr/a:solidFill/a:srgbClr'):
                node.set('val',color)
    def icon(name,kind,tone):
        shape=get(slide,name);path=ASSETS/f'{kind}-{tone}.png'
        if shape is None or not path.is_file():return
        tree=shape._element.getparent();position=tree.index(shape._element)
        new=slide.shapes.add_picture(str(path),shape.left,shape.top,shape.width,shape.height)
        new.name=name
        tree.remove(shape._element);tree.remove(new._element);tree.insert(position,new._element)
    # Text defaults to charcoal; the small section heading retains brand red.
    for shape in slide.shapes:
        if shape.has_text_frame and shape.name!='editorial-kicker':
            for node in shape._element.xpath('.//a:rPr/a:solidFill/a:srgbClr | .//a:defRPr/a:solidFill/a:srgbClr'):
                if node.get('val')==RED:node.set('val',DARK)
    if index==6:
        ink('kpi-decision','AE2349')
        for n,color,kind,tone in ((0,OCEAN,'users','ocean'),(1,JADE,'coins','jade')):
            ink(f'kpi-impact-value-{n}',color)
            icon(f'kpi-impact-icon-{n}',kind,tone)
        for shape in slide.shapes:
            if shape.name.startswith('kpi-to-'):
                recolor_xml(shape._element,{TEAL:'98A2AD'})
    elif index==8:
        for n,color in enumerate((OCEAN,JADE,VIOLET,'AE2349')):
            shape=get(slide,f'flow-node-{n}')
            if shape is not None:rounded(shape,color,50000)
    elif index==11:
        roles=(('pipeline-evidence-icon','map-pinned','ocean'),
               ('pipeline-case-icon','search-check','violet'),
               ('pipeline-click-icon','mouse-pointer-click','slate'),
               ('pipeline-icon-652','git-compare-arrows','ocean'),
               ('pipeline-icon-874','notebook-pen','ochre'),
               ('pipeline-icon-1096','shield-check','jade'),
               ('pipeline-output-icon','files','slate'),
               ('pipeline-data-icon','database','ocean'))
        for args in roles:icon(*args)
        for suffix,color in (('652',OCEAN),('874',OCHRE),('1096',JADE)):
            ink(f'pipeline-model-{suffix}',color)
        for shape in slide.shapes:
            if shape.shape_type==9:recolor_xml(shape._element,{TEAL:'98A2AD'})
    if index in (7,9):
        for shape in slide.shapes:
            if not shape.has_table:continue
            header=SLATE if index==7 else (OCEAN if 'visitors' in shape.name else JADE)
            tint='F0F2F4' if index==7 else ('EDF3F8' if 'visitors' in shape.name else 'EDF5F3')
            for ri,row in enumerate(shape.table.rows):
                for cell in row.cells:
                    cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(header if ri==0 else (WHITE if ri%2 else tint))
        if index==7:ink('adaptation-label',JADE)
        else:
            ink('kpi-total-0',OCEAN);ink('kpi-total-1',JADE)


def reference_case_typography(slide):
    """Keep reference fonts, with readable type on the larger 16:9 canvas."""
    delete(slide,{'case-intro'})
    for n in range(3):
        x=106+n*472
        for shape in slide.shapes:
            if shape.name.startswith('case-') and shape.name.endswith(f'-{n}'):
                shape.top-=42*EMU
        box=get(slide,f'case-panel-{n}')
        if box is not None:box.height=655*EMU
        for key,y,h in (('case-title',303,108),('case-body',428,143),('case-image-caption',772,58),('case-source',840,26)):
            shape=get(slide,f'{key}-{n}')
            if shape is not None:pos(shape,x+24,y,396,h)
    for shape in slide.shapes:
        if not shape.has_text_frame:continue
        name=shape.name
        size=26 if name=='title' else (11 if name=='case-intro' else None)
        if name.startswith(('case-title-','case-badge-')):size=17
        elif name.startswith('case-number-'):size=20
        elif name.startswith('case-body-'):size=15
        elif name.startswith(('case-image-caption-','case-source-')):size=10.5
        elif name in ('editorial-kicker','editorial-page'):size=10
        if size is None:continue
        if name.startswith(('case-title-','case-body-')):
            value=' '.join(shape.text.split())
            fit_text(shape,value,size/0.75,DARK if 'title' in name else MUTED,'title' in name)
        for p in shape.text_frame.paragraphs:
            for run in p.runs:
                # Preserve the size just fitted to long case prose/title. A
                # second forced enlargement invalidates wrapping and overlaps
                # neighbouring cards even though the text box itself fits.
                if not name.startswith(('case-title-','case-body-')):
                    run.font.size=Pt(size)
                run.font.name='Century Gothic'
                props=run._r.get_or_add_rPr()
                east=props.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ea')
                if east is None:east=OxmlElement('a:ea');props.append(east)
                east.set('typeface','맑은 고딕')


def application_photos(slide,report,excluded_blobs=()):
    """Illustrative local facilities, never presented as confirmed program partners."""
    import json
    from .proposal_cover_photo import choose_cover_photo
    body=get(slide,'adaptation-body')
    if body is None:return
    sources=list(report.get('evidence_sources') or [])
    code=str(report.get('region_code') or (report.get('ml_analysis') or {}).get('region_code') or '')
    if code.isdigit():
        path=Path(__file__).resolve().parents[1]/'storage/application_photo_sources'/f'{code}.json'
        try:
            cached=json.loads(path.read_text(encoding='utf-8'))
            if isinstance(cached,list):sources.extend(s for s in cached if isinstance(s,dict))
        except (OSError,ValueError):pass
    categories=((('박물관','전시'),('박물관','미술관','전시관')),
                (('문화유산','역사'),('관아','향교','궁','문화유산')),
                (('공연','문화시설'),('공연장','문화회관','아트센터')))
    pool=[s for s in sources if isinstance(s,dict) and any(any(w in body.text for w in intent) and
          any(w in str(s.get('title','')) for w in names) for intent,names in categories)]
    assets=[];excluded=list(excluded_blobs)
    for _ in range(2):
        asset=choose_cover_photo({'region_name':report.get('region_name'),'evidence_sources':pool},
                                excluded_title=' | '.join(a[0]['title'] for a in assets),excluded_blobs=excluded)
        if not asset:break
        assets.append(asset);excluded.append(asset[1])
    if not assets:
        pos(get(slide,'adaptation-panel'),106,680,1388,128)
        pos(body,478,697,962,95)
        return
    pos(get(slide,'selection-case'),106,182,1388,65)
    matrix=get(slide,'selection-matrix')
    pos(matrix,106,265,1388,330)
    # Native tables use row heights even if the outer frame is resized.
    rows=matrix.table.rows
    if len(rows)>1:
        rows[0].height=50*EMU
        for row in list(rows)[1:]:row.height=round(280*EMU/(len(rows)-1))
    pos(get(slide,'adaptation-panel'),106,619,1388,215)
    label=get(slide,'adaptation-label');pos(label,128,634,736,39)
    pos(body,128,684,740,122);fit_text(body,' '.join(body.text.split()),23,DARK)
    for i,(source,blob) in enumerate(assets):
        x=917+i*278
        pic=slide.shapes.add_picture(BytesIO(blob),round(x*EMU),645*EMU,252*EMU,126*EMU)
        pic.name=f'adaptation-photo-{i}'
        iw,ih=pic.image.size;ratio=iw/ih
        if ratio>2:pic.crop_left=pic.crop_right=(1-2/ratio)/2
        else:pic.crop_top=pic.crop_bottom=(1-ratio/2)/2
        cap=text(slide,f'adaptation-photo-label-{i}',f"적용 공간 예시 · {source['title']}\n사진: 한국관광공사",x,778,252,48,14,MUTED)
        cap.click_action.hyperlink.address=source.get('source_url','')
    slide.notes_slide.notes_text_frame.text+='\n적용 공간 참고 사진(사업 참여 확정 아님): '+json.dumps([a[0] for a in assets],ensure_ascii=False)


def tighter_case_result(slide):
    for shape in slide.shapes:
        if shape.name.startswith('result-'):shape.top-=32*EMU
    heading=get(slide,'result-title')
    if heading:pos(heading,106,162,1388,58)
    scope=get(slide,'result-scope')
    if scope:gray_intro(scope,226)
    photo=get(slide,'result-photo')
    caption=get(slide,'result-photo-source')
    if photo is not None and caption is not None:
        from .proposal_layout_v7 import wrap_words
        value=caption.text
        address=next((r.hyperlink.address for p in caption.text_frame.paragraphs
                      for r in p.runs if r.hyperlink.address), None)
        x=photo.left/EMU+12; width=photo.width/EMU-24
        lines=wrap_words(value,width-24,17)
        height=max(42,len(lines)*22+16)
        y=(photo.top+photo.height)/EMU-height-12
        delete(slide,{'result-photo-source'})
        rect(slide,'result-photo-source-bg',x,y,width,height,CAPTION_BLUE)
        caption=text(slide,'result-photo-source',value,x+12,y+4,width-24,height-8,17,WHITE)
        center(caption)
        if address:caption.click_action.hyperlink.address=address


def gray_intro(shape, y=176):
    """The approved gray lead-in is 26 points; body/caption styles stay separate."""
    pos(shape,106,y,1388,62)
    fit_text(shape,shape.text,26/.75,MUTED)
    center(shape)
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:run.font.size=Pt(26)


def _goal_growth_label(gap, baseline, *, money=False):
    """Choose units per metric before rounding the displayed percentage."""
    rate=gap/baseline*100
    if abs(rate)>=1:
        return f'{rate:+.2f}%'
    amount=round(abs(gap))
    sign='+' if gap>=0 else '−'
    if not money:
        return f'{sign}{amount:,}명'
    if amount<10000:
        return f'{sign}{amount:,}원'
    # Round only for the subtitle; charts and the underlying totals stay exact.
    man=round(amount/10000)
    eok, remainder=divmod(man,10000)
    parts=[]
    if eok:parts.append(f'{eok:,}억')
    if remainder:parts.append(f'{remainder:,}만')
    return f"약 {sign}{' '.join(parts)} 원"


def business_goal_intro(slide, report):
    from .proposal_presentation_v3 import _scenario_rows
    scenario=_scenario_rows(report)
    value='사업기간의 방문자·관광소비액 목표를 월별 전망과 함께 확인합니다.'
    if scenario and scenario['has_target']:
        visitors=sum(scenario['baseline_visitors'])
        spending=sum(scenario['baseline_spending'])
        if visitors>0 and spending>0:
            visitor_label=_goal_growth_label(scenario['visitor_gap'],visitors)
            spending_label=_goal_growth_label(scenario['spending_gap'],spending,money=True)
            value=(f"{scenario['horizon']}개월 목표: ML 전망 대비 방문자 {visitor_label}, "
                   f"관광소비액 {spending_label}")
    intro=text(slide,'business-goal-intro',value,106,176,1388,62,26/.75,MUTED)
    gray_intro(intro)
    for i in range(3):
        x=106+i*470
        block=get(slide,f'block-{i}')
        if block is not None:pos(block,x,257,448,170)
        heading=get(slide,f'block-title-{i}')
        if heading is not None:pos(heading,x+18,272,414,37)
        body=get(slide,f'block-body-{i}')
        if body is not None:
            pos(body,x+18,316,414,104)
            fit_text(body,body.text,23,DARK)


def tighter_kpi(slide):
    intro=get(slide,'kpi-intro')
    if intro:gray_intro(intro)
    for shape in slide.shapes:
        if shape.name.startswith('kpi-') and shape.name not in ('kpi-intro','kpi-footnote'):
            shape.top-=18*EMU


def apply_theme(prs, report):
    """Apply the approved visual theme after population/order, without reauthoring facts."""
    photo_shape=get(prs.slides[2],'visitor-photo')
    photo=photo_shape.image.blob if photo_shape is not None and hasattr(photo_shape,'image') else None
    photo_caption=get(prs.slides[2],'hero-source-label')
    photo_credit=photo_caption.text if photo_caption is not None else ''
    palette={'004EA2':RED,'0054AA':RED,'00B7C9':TEAL,'FF6B00':RED,'008795':RED,
             'E6F4F7':PALE,'F4F7FA':GRAY,'1D2227':DARK,'59636E':MUTED,'D4DFE8':'E4E4E4',
             'DCE5EF':'E4E4E4','0D3158':RED,'003968':RED}
    reserved_photo_blobs=[shape.image.blob for page in prs.slides for shape in page.shapes
                          if hasattr(shape,'image') and ('photo' in shape.name.lower() or 'image' in shape.name.lower())]
    for i,slide in enumerate(prs.slides):
        if i in (0,len(prs.slides)-1):continue
        slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string(GRAY)
        recolor_xml(slide._element,palette)
        delete(slide,{'title-rule','editorial-accent'} | {s.name for s in slide.shapes if s.name.endswith('-accent') or s.name.startswith('block-rail-')})
        title=get(slide,'title')
        if title:
            pos(title,561 if i==2 else 106,title.top/EMU,950 if i==2 else 1388)
            fit_text(title,title.text,40,DARK,True)
            for p in title.text_frame.paragraphs:p.alignment=PP_ALIGN.LEFT
        kicker=get(slide,'editorial-kicker')
        if kicker:
            pos(kicker,561 if i==2 else 106,35,680,43);fit_text(kicker,kicker.text,23,RED,True)
        for shape in list(slide.shapes):
            name=shape.name
            if name.startswith(('case-panel-','execution-panel-','chart-panel-','kpi-impact-panel-')) or name in ('research-background','pipeline-data-band','adaptation-panel') or name in ('block-0','block-1','block-2'):
                rounded(shape,WHITE)
            if name in ('kpi-decision-band','result-change-band'):rounded(shape,PALE)
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.name='Noto Sans KR'
                        if run.hyperlink.address:
                            run.font.color.rgb=RGBColor.from_string(MUTED);run.font.underline=False
                if name in ('case-intro','pipeline-intro','kpi-intro','result-title','result-scope'):
                    center(shape)
                if name=='comparison-note':
                    value=shape.text.replace('파란선','청록선').replace('주황선','빨간선')
                    fit_text(shape,value,19,MUTED)
            if name.startswith('comparison-chart-') and '-legend-line-' in name:
                shape.fill.fore_color.rgb=RGBColor.from_string(TEAL if name.endswith('-0') else RED)
        charts(slide);tables(slide);recolor_icons(slide)
        # White capsule headers like the reference course and process cards.
        if i==3:
            for n in range(3):
                box=get(slide,f'case-panel-{n}')
                if box is None:continue
                box.height=630*EMU
                badge=get(slide,f'case-badge-{n}'); badge.fill.solid();badge.fill.fore_color.rgb=RGBColor.from_string(RED if n==0 else 'FBE2E9')
                fit_text(badge,badge.text,25,WHITE if n==0 else RED,True);center(badge)
                rounded(badge,RED if n==0 else 'FBE2E9',40000)
        if i==4:
            panel(slide,'result-evidence-card',725,330,790,458)
        if i==6:
            panel(slide,'kpi-reason-card',106,285,710,523)
        if i==8:
            for n in range(4):
                step=get(slide,f'flow-node-{n}')
                rounded(step,RED,50000);fit_text(step,step.text,29,WHITE,True);center(step)
        if i==11:
            for x in (652,874,1096):panel(slide,f'pipeline-agent-card-{x}',x,282,180,353)
        if i in (6,7,8,9,11):balanced_colors(slide,i)
        if i==3:reference_case_typography(slide)
        if i==4:tighter_case_result(slide)
        if i==5:business_goal_intro(slide,report)
        if i==6:tighter_kpi(slide)
        if i==7:application_photos(slide,report,reserved_photo_blobs)
        if i==11 and get(slide,'pipeline-intro') is not None:gray_intro(get(slide,'pipeline-intro'),190)
    contents(prs.slides[1]);growth(prs.slides[2])
    from .proposal_cover_photo import choose_cover_photo
    document_photo_blobs=[shape.image.blob for page in list(prs.slides)[1:-1] for shape in page.shapes
                          if hasattr(shape,'image') and ('photo' in shape.name.lower() or 'image' in shape.name.lower())]
    cover_asset=choose_cover_photo(report,photo,photo_credit,excluded_blobs=document_photo_blobs)
    cover(prs.slides[0],cover_asset[1] if cover_asset else None,
          region_name=str(report.get('region_name') or ''),photo_source=cover_asset[0] if cover_asset else None)
    ending_asset=choose_cover_photo(
        report, photo, photo_credit+' | '+(cover_asset[0]['title'] if cover_asset else ''),
        excluded_blobs=tuple(document_photo_blobs)+((cover_asset[1],) if cover_asset else ()))
    cover(prs.slides[-1],ending_asset[1] if ending_asset else None,True,
          region_name=str(report.get('region_name') or ''),photo_source=ending_asset[0] if ending_asset else None)

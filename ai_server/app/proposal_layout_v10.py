"""Evidence-first presentation finishing pass, with editable native content."""
import json
from pathlib import Path
from .proposal_layout_v9 import *
from .case_recommendation import report_cases, case_reference_role, VERSION

ASSETS = Path(__file__).resolve().parents[1] / 'assets' / 'proposal_cases'


def case_narrative(value):
    # Web citations retain their actual URL in the source link and speaker notes.
    # A raw Markdown URL is not prose and can force case cards into tiny type.
    return re.sub(r'\s*\(?\[[^\]]+\]\(https?://[^\s]+\)\)?', '', str(value or '')).strip()


def case_card_summary(source):
    """Display observed festival figures, retaining the event/period scope."""
    if source.get('festival_statistics'):
        return str(source.get('observed_result') or '') + ' 축제 개최기간의 방문 실적을 참고합니다.'
    return case_narrative(source.get('operating_model'))


def safe_text(slide,name,value,x,y,w,h,size=22,color=SLATE,bold=False):
    shape=text(slide,name,value,x,y,w-40,h,size,color,bold)
    shape.width=w*EMU
    return shape


def link(shape, url):
    if str(url or '').startswith(('https://', 'http://')):
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs: run.hyperlink.address = url


def case_cards(slide, report):
    from .case_images import case_image
    from .proposal_case_outcomes import source_outcome
    from PIL import Image
    header(slide, '지역별 참고 사례')
    rows, primary = report_cases(report)
    image_sources=[]
    text(slide,'case-intro','운영 방식이 맞는 공식 사례를 연결합니다. 지역 자체의 유사도 순위와는 구분합니다.',106,188,1388,47,22,SLATE)
    for i, source in enumerate(rows):
        x=106+i*472
        rect(slide,f'case-panel-{i}',x,258,444,542,PALE if source in primary else WHITE)
        text(slide,f'case-number-{i}',f'{i+1:02d}',x+24,277,64,44,29,CYAN,True)
        chosen=source in primary
        text(slide,f'case-badge-{i}','적용 사례' if chosen else '후보 사례',x+99,278,248,44,29,BLUE,True)
        if chosen:
            tick=slide.shapes.add_shape(MSO_SHAPE.OVAL,(x+365)*EMU,282*EMU,43*EMU,43*EMU)
            tick.name=f'case-selected-tick-{i}';tick.fill.solid();tick.fill.fore_color.rgb=RGBColor.from_string(BLUE);tick.line.fill.background()
            mark=text(slide,f'case-selected-mark-{i}','✓',x+365,281,43,43,26,WHITE,True)
            from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
            mark.text_frame.paragraphs[0].alignment=PP_ALIGN.CENTER;mark.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE
        name=(source.get('case_region') or '')+' · '+(source.get('intervention') or source.get('title') or '')
        safe_text(slide,f'case-title-{i}',name,x+24,339,396,98,25,INK,True)
        safe_text(slide,f'case-body-{i}',case_card_summary(source),x+24,452,396,158,21,SLATE)
        outcome=source_outcome(source)
        image_info=(outcome or {}).get('photo') or case_image(source)
        if image_info:
            asset=image_info['path']
            with Image.open(asset) as raw:
                scale=min(396/raw.width,180/raw.height)
                w,h=raw.width*scale,raw.height*scale
            picture=slide.shapes.add_picture(str(asset),int((x+24+(396-w)/2)*EMU),int((615+(180-h)/2)*EMU),int(w*EMU),int(h*EMU))
            picture.name=f'case-photo-{i}'
            link_caption=text(slide,f'case-image-caption-{i}',image_info['caption']+'\n출처: '+image_info['credit'],x+24,805,396,53,16,SLATE)
            link(link_caption,image_info['page_url'])
            image_sources.append({k:str(v) for k,v in image_info.items() if k!='path'})
        else:
            text(slide,f'case-image-caption-{i}','해당 사례의 확인된 이미지 미등록\n아래 공식 사업 원문에서 확인할 수 있습니다.',x+24,705,396,80,18,SLATE)
        shape=text(slide,f'case-source-{i}','사업 원문 보기 ↗',x+24,868,396,24,16,BLUE)
        link(shape,source.get('source_url'))
    slide.notes_slide.notes_text_frame.text=json.dumps({'method':VERSION,'sources':rows,
        'images':image_sources,'image_policy':'Official web assets; exact case or clearly labelled similar operation / general tourism reference. No generated fallback.'},ensure_ascii=False,indent=2)


def selection(slide, report):
    header(slide,'이 사업을 참고한 이유와 지역 적용 방법')
    rows, primary=report_cases(report)
    source=primary[0] if primary else {}
    name=' · '.join(filter(None,[source.get('case_region'),source.get('intervention')]))
    text(slide,'selection-case',name or '동일 운영 방식의 공식 사례를 추가 연결할 수 있습니다.',106,198,1388,70,30,BLUE,True)
    facts=[f"{f.get('metric')}: {f.get('value')}" for f in report.get('observed_findings') or [] if f.get('metric') and f.get('value')]
    local=' / '.join(facts[:3]) or '저장된 지역 관측값이 없습니다.'
    strategy=base._first_strategy(report)
    comparison=[['비교 기준','확인한 근거','이번 기획에 연결한 판단'],
        ['사업의 이용 흐름',case_narrative(source.get('operating_model')) or '동일 방식 문서 미확보','결제·혜택·재이용 또는 예약·체류 등 실제 이용 흐름이 같은 사례만 연결'],
        ['선택 지역의 소비·체류',local,'지역 소비·체류 지표를 보고 참여 업종과 이용 대상을 정합니다.'],
        ['시범 운영 규모',(f"추가 방문 목표 {report['reference_estimate'].get('additional_visitors_target',0):,.0f}명 × 1%\n시범 참여 {report['reference_estimate']['quantity']:,}건(올림·계획 가정)" if report['reference_estimate'].get('additional_visitors_target') and 'quantity' in report['reference_estimate'] else '입력 예산 총액 안에서 참여량과 운영 인력을 배분'), '참여량 × 건별 혜택과 운영 인일로 견적을 계산합니다.']]
    decision = report.get('planning_decision') or {}
    candidates = decision.get('design_candidates') or []
    alternatives = [c.get('title', '') for c in candidates if c.get('candidate_id') != decision.get('selected_candidate_id')]
    if decision.get('selection_reason'):
        comparison[3] = ['후보 비교·선정 이유', '비교 후보: ' + ' / '.join(alternatives) if alternatives else '저장된 후보 비교', case_narrative(decision['selection_reason'])]
    table(slide,'selection-matrix',comparison,106,295,1388,360,[232,578,578],21)
    rect(slide,'adaptation-panel',106,680,1388,155,PALE)
    text(slide,'adaptation-label','사례 → 우리 지역의 제안',128,703,320,65,23,BLUE,True)
    safe_text(slide,'adaptation-body',strategy.get('solution') or strategy.get('title'),478,697,962,124,21,INK)
    label=text(slide,'selection-source','문서 기반 함수 연결 · '+str(source.get('title') or '기록 없음'),106,848,1388,38,18,BLUE)
    link(label,source.get('source_url'))
    # Details remain in speaker notes, outside the presentation narrative.
    # text(slide,'selection-limit','선정 근거: 운영 방식 일치. 인구·자연환경·교통이 같은 지역이라는 뜻이 아니며, 환경 유사성을 입증한 비교는 아닙니다.',106,862,1388,30,16,SLATE)
    slide.notes_slide.notes_text_frame.text=json.dumps({'method':VERSION,'selected_source':source,'observed_findings':report.get('observed_findings'),
        'interpretation':'7개 예측값으로 학습한 추천 모델 없음. 관측 지표는 지역 설계 입력이고 공식 운영 문서가 사례 연결 기준.'},ensure_ascii=False,indent=2)


def merged_execution(slide,report):
    header(slide,'4단계 실행 가이드 예시안')
    groups=execution_groups(report)
    for i,(lead,body) in enumerate(execution_copy(report)):
        y=220+i*143
        rect(slide,f'execution-panel-{i}',106,y,1388,126,WHITE)
        text(slide,f'flow-node-{i}',f'{i+1:02d}',128,y+22,75,51,34,CYAN,True)
        text(slide,f'execution-stage-{i}',LABELS[i],232,y+13,286,41,25,BLUE,True)
        text(slide,f'execution-period-{i}',period(groups[i])+' · '+ROLES[i],232,y+62,286,50,18,SLATE)
        text(slide,f'execution-lead-{i}',lead,558,y+10,906,37,23,INK,True)
        text(slide,f'execution-body-{i}',body,558,y+54,906,64,20,SLATE)
    text(slide,'execution-origin','작성 기준: 저장된 일정 + 사업 유형별 운영 안내 함수. 공식 협약·법정 절차를 인용한 내용이 아니라 조정 가능한 실행 예시입니다.',106,822,1388,60,19,SLATE)
    slide.notes_slide.notes_text_frame.text='proposal_layout_v9.execution_groups / execution_copy → 4개 단계. 일정은 저장 implementation_steps, 업무 설명은 환급형/일반형 코드 템플릿.\n'+json.dumps(base._first_strategy(report).get('implementation_steps'),ensure_ascii=False)


def estimate(slide,report):
    header(slide,'견적 예시안')
    e=report['reference_estimate']
    text(slide,'estimate-label','참여 목표에 따른 예상 사업비',106,203,890,50,29,SLATE,True)
    text(slide,'estimate-total',f"총 {e['total_krw']/10000:,.0f}만 원",1010,197,484,62,39,BLUE,True)

    rows=[['항목','금액','산출근거']]
    for i,item in enumerate(e['items']):
        rows.append([item['name'],f"{item['amount']:,}원",item['basis']])
    sh=table(slide,'estimate-table',rows,106,291,1388,402,[290,280,818],21)
    sh.table.cell(0,2).fill.solid();sh.table.cell(0,2).fill.fore_color.rgb=RGBColor.from_string('16845B')
    note=e.get('scale_basis', '참여량은 운영 기간과 예산에 따른 계획 가정입니다.')+'\n직접비 = 혜택 + 운영·정산 + 시스템 + 홍보 + 평가; 예비비 = 직접비 × 10%. 단가와 처리량은 기획 가정입니다.'
    if 'quantity' not in e:
        note='현재 금액은 사용자가 지정한 총액을 기존 항목 비중으로 재배분한 가정입니다. 각 항목 수량·단가는 별도 조정합니다.'
    elif e.get('scenario_note'):
        note=(e['scenario_note'] + '\n'
              f"100% 참여 시 참고예산 {e['full_participation_budget_krw']:,}원 · 같은 계획 결제·단가 기준, 지급 상한 총액은 아님. 예비비는 직접비의 10%.")
    text(slide,'estimate-method',note,106,720,1388,106,19,SLATE)
    text(slide,'estimate-note','예상 견적이며 실제 액수와 다를 수 있습니다. 세금 포함 여부·업무 범위·수량·단가는 계약 전에 조정합니다.',106,848,1388,38,18,BLUE)
    slide.notes_slide.notes_text_frame.text='operating_target.build_operating_target; 사용자 재배분: idea_proposal.scale_estimate; 전망 없는 기존 견적: proposal_evidence.build_reference_estimate.\n'+json.dumps(e,ensure_ascii=False,indent=2)


def polish(prs,report):
    project_targets(prs.slides[2], report)
    case_cards(prs.slides[4],report)
    selection(prs.slides[5],report)
    merged_execution(prs.slides[6],report)
    # Old slide 8 is fully consolidated into slide 7; no saved report is mutated.
    sid=prs.slides._sldIdLst[7];prs.part.drop_rel(sid.rId);prs.slides._sldIdLst.remove(sid)
    estimate(prs.slides[9],report)
    basis=prs.slides[8]
    from .proposal_result_slide import case_performance
    case_performance(basis,report)
    calculation_pages(prs, report, after=9)
    sid=next(sid for sid in prs.slides._sldIdLst if sid.id==basis.slide_id)
    prs.slides._sldIdLst.remove(sid);prs.slides._sldIdLst.insert(5,sid)
    # User-edited order: purpose, cases, result, goal, goal basis, application.
    requested=['지역별 참고 사례','사례 실적','사업 목표','목표 KPI 산출근거',
               '이 사업을 참고한 이유와 지역 적용 방법','4단계 실행 가이드 예시안',
               '머신러닝 예측값과 목표 KPI','견적 예시안','기획서 생성 파이프라인']
    for index,title in enumerate(requested,3):
        page=next(s for s in prs.slides if any(sh.name=='title' and (sh.text==title or (title=='사례 실적' and sh.text=='참고 사례 운영 방식')) for sh in s.shapes if sh.has_text_frame))
        sid=next(sid for sid in prs.slides._sldIdLst if sid.id==page.slide_id)
        prs.slides._sldIdLst.remove(sid);prs.slides._sldIdLst.insert(index,sid)
    from .proposal_growth_slide import growth_page, section_labels
    growth_page(prs.slides[2],report)
    labels=[('0','개요'),('1','지역별 사례와 실적'),('2','사업 목표와 KPI 산출근거'),
            ('3','지역 적용과 실행'),('4','견적 예시안'),('5','기획서 생성 과정'),('6','근거·데이터·출처')]
    toc=prs.slides[1]
    for shape in list(toc.shapes):
        if shape.name.startswith('toc-'):shape._element.getparent().remove(shape._element)
    for i,(number,label) in enumerate(labels):
        col,row=divmod(i,4);x=562+540*col;y=296+108*row
        text(toc,f'toc-{i}-number',number,x,y,61,44,25,CYAN,True)
        text(toc,f'toc-{i}-label',label,x+65,y+5,415,58,25,INK,True)
    from .proposal_visual_finish import finish_design
    finish_design(prs,report)
    section_labels(prs)
    from .proposal_theme_spicus import apply_theme
    apply_theme(prs, report)


def project_targets(slide, report):
    from .proposal_calculation_basis import calculation_basis
    b=calculation_basis(report)
    for shape in list(slide.shapes):
        if shape.name in ('s3-operation', 's3-operation-rule'):
            shape._element.getparent().remove(shape._element)
    text(slide,'s3-goal-label','사업의 목표 · 3개월 월별 합계 기준',561,645,895,40,23,BLUE,True)
    for i,r in enumerate(b['totals']):
        x=561+i*459
        rect(slide,f's3-goal-panel-{i}',x,698,438,132,PALE)
        text(slide,f's3-goal-name-{i}',r['label'],x+15,709,400,30,21,BLUE,True)
        delta=(f"{r['additional']:,.0f}명" if r['key']=='visitors' else f"{r['additional']/1e8:,.2f}억 원")
        text(slide,f's3-goal-value-{i}',f'↑ {delta}',x+15,746,404,41,30,ORANGE,True)
        pct=f"+{r['pct']:.2f}%" if r['pct'] is not None else '비율 산정 제외'
        text(slide,f's3-goal-pct-{i}',f'기준 전망 대비 {pct} · 계획 목표',x+15,792,404,30,18,SLATE)
    if not b['totals']:
        text(slide,'s3-goal-missing','저장된 사업기간 전망과 목표율을 연결하면 방문·소비 목표를 표시합니다.',561,734,895,100,23,SLATE)


def calculation_pages(prs, report, after):
    from .proposal_calculation_basis import calculation_basis, methodology_sections
    from pptx.opc.packuri import PackURI
    # Earlier passes remove/reorder template slides. Renumber existing parts
    # before add_slide chooses len(slides)+1 to avoid duplicate ZIP part names.
    for index, slide in enumerate(prs.slides, 1):
        slide.part._partname=PackURI(f'/ppt/slides/slide{index}.xml')
    b=calculation_basis(report);sections=methodology_sections(report)
    from .proposal_infographics import calculations, ml_to_plan, planning_rationale_page
    from .planning_rationale import report_rationale, rationale_sections
    first=prs.slides.add_slide(prs.slides[-1].slide_layout)
    calculations(first,report)
    second=prs.slides.add_slide(prs.slides[-1].slide_layout)
    ml_to_plan(second,report)
    for slide in (first,second):
        script=(
            '발표 설명\n방문·소비의 기준 전망은 저장 모델에서 나온 수치입니다. 운영 규모와 참여 가정에서 추가 방문·소비를 계산합니다. '
            '지역 전체 사례 증가율이나 인구 비율을 신규 사업의 효과로 복사하지 않습니다. 운영량과 단가를 곱해 예상 견적을 산정합니다. '
            '자동 제안과 사용자 수정 목표는 구분하며 계획 시나리오는 통계적 신뢰구간이 아닙니다.\n\n'
            if slide==first else
            '발표 설명\n우리 ML은 7개 관광지표의 전망을 계산합니다. 방문자와 소비액은 그래프·목표 규모를 계산하는 기준선입니다. '
            '체류·숙박·검색 관련 5개 전망은 지역의 관광 흐름을 검토할 입력입니다. 서버가 이 수치와 모델 평가를 공식 사례와 함께 Qwen에 제공합니다. '
            'Qwen은 지역 조건과 사례의 운영 방식을 비교하고 Gemma는 선택안을 본문으로 구성합니다. 따라서 ML은 수치 전망을 담당하고 사업 선택은 근거 문서와 LLM 판단을 결합합니다. '
            '지표별로 전년 동월 기준모델이 선택될 수 있으며, 이를 고도화 모델의 성능 개선이라고 설명하지 않습니다.\n\n')
        slide.notes_slide.notes_text_frame.text=script+json.dumps({'sections':sections,'basis':b,
            'icons':{'library':'Lucide 0.468.0','source':'https://github.com/lucide-icons/lucide/tree/0.468.0',
                     'license':'ISC / Feather MIT; ai_server/assets/proposal_icons/LICENSE'},
            'pipeline':{'agents':['Evidence','Case Scout','Transferability','Planner','Reviewer'],
                        'implementation':'app/agents/report_orchestrator.py; Evidence + Case Scout gathered together',
                        'provider_labels':'Completed provider records in this report, otherwise saved evidence / configured model'},
            'functions':['planning_evidence.build_planning_ml_evidence','case_recommendation.link_decision',
                         'report_projection.execution_target','proposal_evidence.build_reference_estimate']},ensure_ascii=False,indent=2)
    pages = [first, second]
    if report_rationale(report):
        third = prs.slides.add_slide(prs.slides[-1].slide_layout)
        planning_rationale_page(third, report)
        third.notes_slide.notes_text_frame.text = (
            '발표 설명\n관측값과 ML 전망을 먼저 확인하고, 사례의 운영 방식에서 가져온 지역 적용 가설을 구분했습니다. '
            '후보마다 이용 행동·운영 준비·성과 확인이라는 같은 기준을 사용합니다. 이 선택 이유는 조건부 기획 판단이며 성과 보장이 아닙니다.\n\n' +
            '\n\n'.join(title+'\n'+body for title,body in rationale_sections(report)) + '\n\n' +
            json.dumps({'decision': report.get('planning_decision'), 'sources': report.get('evidence_sources')}, ensure_ascii=False, indent=2))
        pages.append(third)
    # New traced decisions include their rationale; older reports keep the approved pages.
    for offset,slide in enumerate(pages,1):
        sid=next(sid for sid in prs.slides._sldIdLst if sid.id==slide.slide_id)
        prs.slides._sldIdLst.remove(sid);prs.slides._sldIdLst.insert(after+offset,sid)

import unittest
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from ai_server.test_proposal_presentation_v3_contract import _sample_report, _slide_text
from ai_server.app.proposal_presentation_v4 import create_strategy_proposal_presentation
from ai_server.app.proposal_layout_v9 import execution_copy, case_selection_basis, execution_groups, period
from ai_server.presentation_test_support import slide_with_title

class ExecutionDetailTest(unittest.TestCase):
    def test_last_word_case_has_no_orphan_spacer(self):
        from docx import Document
        from docx.oxml.ns import qn
        from unittest.mock import patch
        from ai_server.app.proposal_document_v2 import case_card
        doc = Document()
        source = {'intervention': '공식 체험 사업', 'operating_model': '예약 후 체험합니다.'}
        with patch('ai_server.app.proposal_document_v2.case_image', return_value=None):
            case_card(doc, {}, source, 0, [], add_gap=False)
        self.assertEqual(list(doc._element.body)[-2].tag, qn('w:tbl'))

    def test_case_theme_does_not_enlarge_fitted_long_prose(self):
        from ai_server.app.proposal_layout_v7 import text, EMU, _width
        from ai_server.app.proposal_theme_spicus import reference_case_typography
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        narrative = '공식 운영 자료의 대상과 비용, 참여 절차와 성과 집계 기간을 구분하여 사례의 조건을 확인합니다. ' * 4
        shape = text(slide, 'case-body-0', narrative, 130, 452, 396, 400, 20)
        reference_case_typography(slide)
        self.assertEqual(' '.join(shape.text.split()), narrative.strip())
        for p in shape.text_frame.paragraphs:
            size = p.runs[0].font.size.pt / .75
            self.assertLess(p.runs[0].font.size.pt, 15)
            self.assertLessEqual(_width(p.text, size), shape.width / EMU - 12)

    def test_result_slide_keeps_long_source_url_as_link_not_prose(self):
        from copy import deepcopy
        from unittest.mock import patch
        from ai_server.app.proposal_result_slide import case_performance
        url = 'https://example.go.kr/official?' + 'query=value&' * 100
        narrative = '여행자가 지역에서 지출한 금액 일부를 상품권으로 환급하여 지역 상점에서 다시 사용하도록 운영합니다.'
        source = {'source_id': 'case:refund', 'source_url': url,
                  'case_region': '사례 지역', 'intervention': '지역 환급',
                  'operating_model': narrative + f' ([공식 출처]({url}))'}
        before = deepcopy(source)
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        with patch('ai_server.app.case_recommendation.report_cases', return_value=([source], [source])), \
                patch('ai_server.app.proposal_result_slide.report_outcome', return_value=None), \
                patch('ai_server.app.case_images.case_image', return_value=None):
            case_performance(slide, {})
        shape = next(s for s in slide.shapes if s.name == 'result-operation')
        self.assertEqual(' '.join(shape.text.split()), narrative)
        self.assertTrue(all(r.hyperlink.address == url for p in shape.text_frame.paragraphs for r in p.runs))
        self.assertIn(url, slide.notes_slide.notes_text_frame.text)
        self.assertEqual(source, before)

    def test_word_keeps_capacity_disclosures_together_without_changing_data(self):
        from copy import deepcopy
        from docx import Document
        from docx.shared import Pt
        from ai_server.app.proposal_document import create_strategy_proposal_document
        report = _sample_report()
        before = deepcopy(report)
        doc = Document(create_strategy_proposal_document(report))
        notes = [p for p in doc.paragraphs if '산출한 추가 방문·소비 규모는' in p.text]
        self.assertEqual(len(notes), 1)
        self.assertIn('통계적 오차범위가 아닙니다', notes[0].text)
        self.assertIn('실측 객단가가 아닙니다', notes[0].text)
        self.assertLessEqual(notes[0].paragraph_format.space_after, Pt(4))
        self.assertEqual(notes[0].runs[0].font.size, Pt(9))
        cells = '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        self.assertIn('전국 비교 캐시가 없으면 저장 ML로 비교값을 계산합니다', cells)
        self.assertEqual(report, before)

    def test_five_steps_preserve_operation_and_evaluation_dates_in_both_outputs(self):
        from copy import deepcopy
        from docx import Document
        from ai_server.app.proposal_document import create_strategy_proposal_document
        r = _sample_report()
        r['strategies'][0]['implementation_steps'] = [
            {'step':1, 'schedule':'2026-10', 'task':'참여 시설 조사', 'deliverable':'시설 목록'},
            {'step':2, 'schedule':'2026-10', 'task':'콘텐츠와 안전 계획 수립', 'deliverable':'운영 계획'},
            {'step':3, 'schedule':'2026-10', 'task':'예약 시스템 점검', 'deliverable':'점검 결과'},
            {'step':4, 'schedule':'2026-11 ~ 2026-12', 'task':'프로그램 시범 운영', 'deliverable':'운영 원장'},
            {'step':5, 'schedule':'2026-12', 'task':'성과 분석', 'deliverable':'결과 보고서'},
        ]
        before = deepcopy(r)
        groups = execution_groups(r)
        self.assertEqual([v for g in groups for v in g], r['strategies'][0]['implementation_steps'])
        self.assertEqual(period(groups[2]), '2026-11 ~ 2026-12')
        self.assertEqual(period(groups[3]), '2026-12')
        deck = Presentation(create_strategy_proposal_presentation(r))
        periods = {s.name:s.text for s in slide_with_title(deck,'4.1 4단계 실행 가이드 예시안').shapes if s.name.startswith('execution-period-')}
        self.assertIn('2026-11 ~ 2026-12', periods['execution-period-2'])
        self.assertTrue(periods['execution-period-3'].startswith('2026-12'))
        doc = Document(create_strategy_proposal_document(r))
        cells=[row.cells[0].text for table in doc.tables for row in table.rows]
        pilot=[cell for cell in cells if cell.startswith('03  시범 운영')]
        self.assertEqual(len(pilot),1)
        self.assertIn('2026-11 ~ 2026-12',pilot[0])
        evaluation=[cell for cell in cells if cell.startswith('04  실적 확인과 평가')]
        self.assertEqual(len(evaluation),1)
        self.assertIn('2026-12',evaluation[0])
        self.assertEqual(r, before)

    def test_selection_uses_linked_case_without_environment_claim(self):
        r=_sample_report()
        r['strategies'][0]['title']='여행비 환급'
        r['strategies'][0]['solution']='지역상품권 환급'
        r['planning_decision']={'selected_candidate_id':'chosen','design_candidates':[
            {'candidate_id':'chosen','title':'관광 프로그램','case_source_ids':['case:chosen'],'mechanism':'지역 상점 재이용'}]}
        r['evidence_sources']=[
            {'source_id':'case:other','source_type':'benchmark_case','title':'다른 지역'},
            {'source_id':'case:chosen','source_type':'benchmark_case','title':'선정 사례','operating_model':'지역상품권으로 돌려받아 다시 이용',
             'source_url':'https://example.go.kr/case'}]
        basis=case_selection_basis(r)
        self.assertEqual(basis['source']['source_id'],'case:chosen')
        deck=Presentation(create_strategy_proposal_presentation(r))
        slide=slide_with_title(deck,'2.2 적용 사례 운영 방식')
        cards=[s for s in slide.shapes if s.name.startswith('case-adaptation-body-')]
        self.assertEqual(len(cards),3)
        self.assertNotIn('지리·교통·관광자원', ' '.join(s.text for s in cards))
        self.assertIn('https://example.go.kr/case',slide.notes_slide.notes_text_frame.text)
        self.assertNotIn('다른 지역', ' '.join(s.text for s in cards))
    def test_page_order_and_overview_headers(self):
        p=Presentation(create_strategy_proposal_presentation(_sample_report()))
        self.assertIn('월별 방문·소비 증가 목표',_slide_text(p.slides[9]))
        self.assertFalse(any(s.has_text_frame and s.text == '3.1 사례 선정과 지역 적용' for slide in p.slides for s in slide.shapes))
        self.assertIn('4단계 실행 가이드 예시안',_slide_text(p.slides[10]))
        self.assertIn('머신러닝 예측값',_slide_text(p.slides[11]))
        self.assertEqual(len([s for s in p.slides[10].shapes if s.name.startswith('flow-node-')]),4)
    def test_refund_instructions_not_for_unrelated_program(self):
        r=_sample_report();r['strategies'][0]['title']='반값 여행 환급'
        self.assertIn('영수증',str(execution_copy(r)))
        r['strategies'][0]['title']='지역 문화 체험';r['strategies'][0]['solution']='문화시설 예약 운영';r['planning_decision']={}
        self.assertNotIn('환급',str(execution_copy(r)))
if __name__=='__main__':unittest.main()
